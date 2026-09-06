from __future__ import annotations

import io
import math
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from .shared import (
    DEFAULT_PROFILE,
    DEFAULT_THEME,
    JsonArgumentParser,
    PRESENTATION_MARKER,
    SkillError,
    SUPPORTED_PROFILES,
    SUPPORTED_THEMES,
    asset_path,
    atomic_save,
    emit,
    optional_text,
    path_result,
    project_path,
    project_root,
    read_json_spec,
    reject_unknown,
    require_pptx,
    required_text,
    run,
    presentation_style,
    validate_image_file,
)


SLIDE_WIDTH_IN = 13.333333
SLIDE_HEIGHT_IN = 7.5
ECNU_RED = (155, 32, 52)
INK = (34, 34, 34)
MUTED = (107, 112, 120)
SOFT = (238, 239, 241)
WHITE = (255, 255, 255)
DIGITAL_NAVY = (24, 38, 56)
DIGITAL_BLUE = (67, 111, 137)
DIGITAL_PALE = (238, 244, 247)
DIGITAL_LINE = (177, 197, 210)
EDITORIAL_WARM = (246, 242, 235)
EDITORIAL_LINE = (191, 181, 168)
FONT_NAME = "Microsoft YaHei"
THEME_PALETTES = {
    "modern-clean": {
        "accent": (61, 88, 138), "ink": (31, 41, 55), "muted": (100, 112, 128), "soft": (239, 243, 248),
        "navy": (37, 54, 84), "blue": (79, 112, 160), "pale": (238, 244, 250), "line": (186, 201, 220),
        "warm": (247, 247, 245), "editorial_line": (196, 198, 202),
    },
    "academic-editorial": {
        "accent": (142, 67, 47), "ink": (45, 41, 38), "muted": (112, 103, 96), "soft": (243, 239, 233),
        "navy": (63, 55, 50), "blue": (126, 91, 72), "pale": (246, 241, 235), "line": (205, 190, 177),
        "warm": (246, 242, 235), "editorial_line": (191, 181, 168),
    },
    "digital-tech": {
        "accent": (0, 126, 167), "ink": (22, 37, 52), "muted": (83, 104, 119), "soft": (233, 243, 247),
        "navy": (15, 43, 63), "blue": (0, 126, 167), "pale": (231, 244, 248), "line": (151, 196, 211),
        "warm": (241, 246, 248), "editorial_line": (174, 204, 214),
    },
    "warm-education": {
        "accent": (197, 102, 47), "ink": (55, 47, 39), "muted": (119, 103, 87), "soft": (249, 242, 230),
        "navy": (74, 82, 58), "blue": (105, 132, 88), "pale": (241, 246, 234), "line": (198, 207, 177),
        "warm": (250, 244, 233), "editorial_line": (210, 190, 164),
    },
}
SUPPORTED_LAYOUTS = {
    "cover",
    "section",
    "bullets",
    "summary",
    "two-column",
    "metrics",
    "timeline",
    "feature-grid",
    "roadmap",
    "image",
    "quote",
    "closing",
}
VISUAL_COMPOSITION_LAYOUTS = {"image", "timeline", "feature-grid", "roadmap"}
STRUCTURAL_LAYOUTS = {"cover", "section", "closing"}


def apply_theme_palette(theme: str) -> None:
    if not theme or theme == "ecnu-liwa":
        return
    palette = THEME_PALETTES[theme]
    global ECNU_RED, INK, MUTED, SOFT, DIGITAL_NAVY, DIGITAL_BLUE, DIGITAL_PALE, DIGITAL_LINE, EDITORIAL_WARM, EDITORIAL_LINE
    ECNU_RED = palette["accent"]
    INK = palette["ink"]
    MUTED = palette["muted"]
    SOFT = palette["soft"]
    DIGITAL_NAVY = palette["navy"]
    DIGITAL_BLUE = palette["blue"]
    DIGITAL_PALE = palette["pale"]
    DIGITAL_LINE = palette["line"]
    EDITORIAL_WARM = palette["warm"]
    EDITORIAL_LINE = palette["editorial_line"]


def parse_args() -> Any:
    parser = JsonArgumentParser(description="Create a themed PPTX from a constrained JSON specification.")
    parser.add_argument("--project-root", required=True)
    parser.add_argument("--output", required=True)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--spec-json")
    group.add_argument("--spec-file")
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def rgb(value: tuple[int, int, int]) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor(*value)


def inches(value: float) -> Any:
    from pptx.util import Inches

    return Inches(value)


def visual_text_units(value: str) -> float:
    total = 0.0
    for character in value:
        width = unicodedata.east_asian_width(character)
        total += 1.0 if width in {"W", "F", "A"} else 0.55
    return total


def points(value: float) -> Any:
    from pptx.util import Pt

    return Pt(value)


def style_run(run: Any, *, size: float, color: tuple[int, int, int], bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = points(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(
    slide: Any,
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    size: float,
    color: tuple[int, int, int] | None = None,
    bold: bool = False,
    align: Any | None = None,
    vertical_anchor: Any | None = None,
) -> Any:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    color = INK if color is None else color
    shape = slide.shapes.add_textbox(inches(x), inches(y), inches(width), inches(height))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = vertical_anchor if vertical_anchor is not None else MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align if align is not None else PP_ALIGN.LEFT
    paragraph.space_before = points(0)
    paragraph.space_after = points(0)
    paragraph.line_spacing = 1.12
    run = paragraph.add_run()
    run.text = text
    style_run(run, size=size, color=color, bold=bold)
    return shape


def add_rule(
    slide: Any,
    *,
    x: float,
    y: float,
    width: float,
    color: tuple[int, int, int] | None = None,
) -> Any:
    from pptx.enum.shapes import MSO_SHAPE

    color = ECNU_RED if color is None else color
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(x), inches(y), inches(width), inches(0.035))
    shape.name = "ECNU Rule"
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_rect(
    slide: Any,
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    fill: tuple[int, int, int],
    line: tuple[int, int, int] | None = None,
    radius: bool = False,
) -> Any:
    from pptx.enum.shapes import MSO_SHAPE

    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, inches(x), inches(y), inches(width), inches(height))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = points(0.8)
    return shape


def add_circle(
    slide: Any,
    *,
    name: str,
    x: float,
    y: float,
    diameter: float,
    fill: tuple[int, int, int],
    line: tuple[int, int, int] | None = None,
) -> Any:
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        inches(x),
        inches(y),
        inches(diameter),
        inches(diameter),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = points(0.8)
    return shape


def add_profile_marker(slide: Any, profile: str) -> None:
    if profile == "academic-editorial":
        add_rect(
            slide,
            name=f"ECNU Profile Marker {{{profile}}}",
            x=10.14,
            y=0.36,
            width=0.055,
            height=0.36,
            fill=ECNU_RED,
        )
    elif profile == "digital-campus":
        add_rect(
            slide,
            name=f"ECNU Profile Marker {{{profile}}}",
            x=9.82,
            y=0.37,
            width=0.36,
            height=0.12,
            fill=DIGITAL_NAVY,
        )
        add_rect(
            slide,
            name="ECNU Digital Signal",
            x=9.82,
            y=0.54,
            width=0.2,
            height=0.05,
            fill=ECNU_RED,
        )
    else:
        add_rect(
            slide,
            name=f"ECNU Profile Marker {{{profile}}}",
            x=10.05,
            y=0.42,
            width=0.12,
            height=0.12,
            fill=ECNU_RED,
        )


def add_base(
    presentation: Any,
    *,
    profile: str,
    theme: str,
    branded: bool,
    layout: str,
    page_number: int,
    footer: str,
    background: Path | None,
    logo_a: Path | None,
    logo_c: Path | None,
) -> Any:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if branded:
        background_shape = slide.shapes.add_picture(
            str(background), 0, 0, width=presentation.slide_width, height=presentation.slide_height,
        )
        background_shape.name = "ECNU Background"
        logo = logo_a if layout == "cover" else logo_c
        logo_size = 0.62 if layout == "cover" else 0.36
        logo_shape = slide.shapes.add_picture(str(logo), inches(0.72), inches(0.34), width=inches(logo_size))
        logo_shape.name = "ECNU Logo Form A" if layout == "cover" else "ECNU Logo Form C"
        add_text(
            slide, name="ECNU Brand", x=1.5 if layout == "cover" else 1.2, y=0.47 if layout == "cover" else 0.42,
            width=3.8 if layout == "cover" else 3.5, height=0.3 if layout == "cover" else 0.24,
            text="EAST CHINA NORMAL UNIVERSITY" if layout == "cover" else "华东师范大学  ECNU",
            size=9 if layout == "cover" else 8.5, color=ECNU_RED, bold=True,
        )
    else:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_profile_marker(slide, profile)
    if theme:
        for shape in slide.shapes:
            if (shape.name or "").startswith("ECNU Profile Marker {"):
                shape.name = f"Presentation Theme Marker {{{theme}}}"
    if footer:
        add_text(
            slide,
            name="ECNU Footer",
            x=0.72,
            y=7.07,
            width=6.6,
            height=0.2,
            text=footer,
            size=7.5,
            color=MUTED,
        )
    add_text(
        slide,
        name="ECNU Page Number",
        x=9.95,
        y=7.03,
        width=0.45,
        height=0.2,
        text=f"{page_number:02d}",
        size=7.5,
        color=MUTED,
        align=1,
    )
    return slide


def add_bullet_list(
    slide: Any,
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    bullets: list[dict[str, Any]],
    normal_color: tuple[int, int, int] | None = None,
    accent_color: tuple[int, int, int] | None = None,
    primary_size: float = 21,
    secondary_size: float = 17,
    item_spacing: float = 11,
) -> Any:
    from pptx.enum.text import MSO_ANCHOR

    normal_color = INK if normal_color is None else normal_color
    accent_color = ECNU_RED if accent_color is None else accent_color
    shape = slide.shapes.add_textbox(inches(x), inches(y), inches(width), inches(height))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        level = item["level"]
        paragraph.level = level
        paragraph.space_before = points(0)
        paragraph.space_after = points(item_spacing if level == 0 else max(5, item_spacing - 4))
        paragraph.line_spacing = 1.12
        prefix = "●  " if level == 0 else "–  "
        run = paragraph.add_run()
        run.text = prefix + item["text"]
        style_run(
            run,
            size=primary_size if level == 0 else secondary_size,
            color=accent_color if item["accent"] else normal_color,
            bold=item["accent"],
        )
    return shape


def parse_bullets(
    value: Any,
    *,
    label: str,
    maximum_items: int,
    maximum_characters: int,
    maximum_total_characters: int,
) -> list[dict[str, Any]]:
    if not isinstance(value, list) or not value:
        raise SkillError("invalid_bullets", f"{label} must be a non-empty array.")
    if len(value) > maximum_items:
        raise SkillError(
            "too_many_bullets",
            f"{label} has too many items for the fixed layout.",
            details={"count": len(value), "maximum": maximum_items},
        )
    result: list[dict[str, Any]] = []
    for index, item in enumerate(value, start=1):
        if isinstance(item, str):
            text = required_text(item, label=f"{label}[{index}]", maximum=maximum_characters)
            result.append({"text": text, "level": 0, "accent": False})
            continue
        if not isinstance(item, dict):
            raise SkillError("invalid_bullet", f"{label}[{index}] must be a string or object.")
        reject_unknown(item, {"text", "level", "accent"}, label=f"{label}[{index}]")
        text = required_text(
            item.get("text"),
            label=f"{label}[{index}].text",
            maximum=maximum_characters,
        )
        level = item.get("level", 0)
        accent = item.get("accent", False)
        if isinstance(level, bool) or level not in (0, 1):
            raise SkillError("invalid_bullet_level", f"{label}[{index}].level must be 0 or 1.")
        if not isinstance(accent, bool):
            raise SkillError("invalid_bullet_accent", f"{label}[{index}].accent must be boolean.")
        result.append({"text": text, "level": level, "accent": accent})
    total_characters = sum(len(item["text"]) for item in result)
    if total_characters > maximum_total_characters:
        raise SkillError(
            "bullets_too_long",
            f"{label} contains too much text for the fixed layout; split it across slides.",
            details={"length": total_characters, "maximum": maximum_total_characters},
        )
    return result


def add_title(slide: Any, title: str, *, kicker: str = "", layout: str) -> None:
    if kicker:
        add_text(
            slide,
            name=f"ECNU Text Kicker [{layout}]",
            x=0.95,
            y=0.98,
            width=9.5,
            height=0.25,
            text=kicker.upper(),
            size=10,
            color=ECNU_RED,
            bold=True,
        )
    add_text(
        slide,
        name=f"ECNU Title [{layout}]",
        x=0.95,
        y=1.24 if kicker else 1.08,
        width=9.55,
        height=0.65,
        text=title,
        size=32,
        color=INK,
        bold=True,
    )
    add_rule(slide, x=0.95, y=1.88 if kicker else 1.72, width=0.72)


def render_cover(slide: Any, spec: dict[str, Any], profile: str, theme: str) -> None:
    reject_unknown(spec, {"layout", "title", "subtitle", "meta"}, label="cover slide")
    title = required_text(spec.get("title"), label="cover.title", maximum=48)
    subtitle = optional_text(spec.get("subtitle"), label="cover.subtitle", maximum=90)
    meta = optional_text(spec.get("meta"), label="cover.meta", maximum=100)
    if profile == "academic-editorial":
        add_rect(
            slide,
            name="ECNU Editorial Cover Field",
            x=0.88,
            y=1.42,
            width=1.42,
            height=4.92,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Text Editorial Label [cover]",
            x=1.13,
            y=1.75,
            width=0.9,
            height=2.5,
            text="LEARNING\nSTORY" if theme == "warm-education" else "ACADEMIC\nSTORY",
            size=11,
            color=WHITE,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [cover]",
            x=2.85,
            y=1.72,
            width=7.25,
            height=1.8,
            text=title,
            size=40,
            color=INK,
            bold=True,
        )
        add_rule(slide, x=2.88, y=3.72, width=2.25)
        if subtitle:
            add_text(
                slide,
                name="ECNU Text Subtitle [cover]",
                x=2.88,
                y=4.02,
                width=6.9,
                height=0.85,
                text=subtitle,
                size=17,
                color=MUTED,
            )
        if meta:
            add_text(
                slide,
                name="ECNU Text Meta [cover]",
                x=2.88,
                y=5.66,
                width=6.9,
                height=0.42,
                text=meta,
                size=10.5,
                color=ECNU_RED,
                bold=True,
            )
        return
    if profile == "digital-campus":
        add_rect(
            slide,
            name="ECNU Digital Cover Rail",
            x=0.86,
            y=1.38,
            width=0.14,
            height=4.98,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Text Digital Label [cover]",
            x=1.34,
            y=1.82,
            width=3.1,
            height=0.28,
            text="DIGITAL SYSTEM" if theme == "digital-tech" else "DIGITAL CAMPUS · ECNU",
            size=10,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [cover]",
            x=1.3,
            y=2.28,
            width=7.35,
            height=1.75,
            text=title,
            size=40,
            color=INK,
            bold=True,
        )
        if subtitle:
            add_text(
                slide,
                name="ECNU Text Subtitle [cover]",
                x=1.32,
                y=4.16,
                width=7.0,
                height=0.65,
                text=subtitle,
                size=20,
                color=MUTED,
            )
        if meta:
            add_text(
                slide,
                name="ECNU Text Meta [cover]",
                x=1.32,
                y=5.55,
                width=6.8,
                height=0.35,
                text=meta,
                size=12,
                color=ECNU_RED,
                bold=True,
            )
        for index in range(4):
            add_rect(
                slide,
                name=f"ECNU Digital Grid {index + 1}",
                x=8.66 + (index % 2) * 0.45,
                y=4.75 + (index // 2) * 0.45,
                width=0.28,
                height=0.28,
                fill=ECNU_RED if index == 0 else (DIGITAL_NAVY if index == 1 else DIGITAL_PALE),
                line=None if index < 2 else DIGITAL_LINE,
            )
        return
    add_rule(slide, x=0.92, y=1.75, width=0.95)
    add_text(
        slide,
        name="ECNU Title [cover]",
        x=0.92,
        y=2.04,
        width=8.9,
        height=1.35,
        text=title,
        size=40,
        color=INK,
        bold=True,
    )
    if subtitle:
        add_text(
            slide,
            name="ECNU Text Subtitle [cover]",
            x=0.95,
            y=3.52,
            width=8.6,
            height=0.65,
            text=subtitle,
            size=18,
            color=MUTED,
        )
    if meta:
        add_text(
            slide,
            name="ECNU Text Meta [cover]",
            x=0.95,
            y=5.7,
            width=8.6,
            height=0.42,
            text=meta,
            size=11,
            color=ECNU_RED,
            bold=True,
        )


def render_section(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "number", "title", "subtitle"}, label="section slide")
    number = optional_text(spec.get("number"), label="section.number", maximum=8)
    title = required_text(spec.get("title"), label="section.title", maximum=42)
    subtitle = optional_text(spec.get("subtitle"), label="section.subtitle", maximum=100)
    if profile == "academic-editorial":
        add_rect(
            slide,
            name="ECNU Editorial Section Number Field",
            x=0.94,
            y=1.55,
            width=2.15,
            height=4.8,
            fill=EDITORIAL_WARM,
            line=EDITORIAL_LINE,
        )
        add_text(
            slide,
            name="ECNU Text Section Number [section]",
            x=1.23,
            y=2.02,
            width=1.55,
            height=1.15,
            text=number or "§",
            size=45,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Text Editorial Chapter [section]",
            x=1.25,
            y=4.95,
            width=1.45,
            height=0.55,
            text="CHAPTER",
            size=8.5,
            color=MUTED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [section]",
            x=3.65,
            y=2.24,
            width=6.3,
            height=1.3,
            text=title,
            size=36,
            color=INK,
            bold=True,
        )
        add_rule(slide, x=3.68, y=3.82, width=1.55)
        if subtitle:
            add_text(
                slide,
                name="ECNU Text Subtitle [section]",
                x=3.68,
                y=4.12,
                width=5.9,
                height=0.8,
                text=subtitle,
                size=15,
                color=MUTED,
            )
        return
    if profile == "digital-campus":
        add_rect(
            slide,
            name="ECNU Digital Section Rail",
            x=0.88,
            y=1.68,
            width=0.12,
            height=4.48,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Text Section Number [section]",
            x=1.35,
            y=2.08,
            width=1.55,
            height=1.2,
            text=number or "00",
            size=54,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [section]",
            x=3.08,
            y=2.32,
            width=6.7,
            height=1.15,
            text=title,
            size=38,
            color=INK,
            bold=True,
        )
        add_rule(slide, x=3.1, y=3.68, width=1.1)
        if subtitle:
            add_text(
                slide,
                name="ECNU Text Subtitle [section]",
                x=3.12,
                y=4.02,
                width=6.45,
                height=0.72,
                text=subtitle,
                size=18,
                color=MUTED,
            )
        add_rect(
            slide,
            name="ECNU Digital Section Signal",
            x=8.65,
            y=2.16,
            width=0.92,
            height=0.1,
            fill=DIGITAL_PALE,
            line=DIGITAL_LINE,
        )
        add_rect(
            slide,
            name="ECNU Digital Section Signal Accent",
            x=9.33,
            y=2.16,
            width=0.24,
            height=0.1,
            fill=ECNU_RED,
        )
        return
    if number:
        add_text(
            slide,
            name="ECNU Text Section Number [section]",
            x=0.95,
            y=1.65,
            width=2.0,
            height=0.9,
            text=number,
            size=42,
            color=ECNU_RED,
            bold=True,
        )
    add_text(
        slide,
        name="ECNU Title [section]",
        x=0.95,
        y=2.7,
        width=8.9,
        height=1.0,
        text=title,
        size=36,
        color=INK,
        bold=True,
    )
    add_rule(slide, x=0.95, y=3.82, width=1.15)
    if subtitle:
        add_text(
            slide,
            name="ECNU Text Subtitle [section]",
            x=0.95,
            y=4.12,
            width=8.9,
            height=0.7,
            text=subtitle,
            size=16,
            color=MUTED,
        )


def render_summary(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "bullets"}, label="summary slide")
    title = required_text(spec.get("title"), label="summary.title", maximum=160)
    bullets = spec.get("bullets", [])
    if not isinstance(bullets, list) or len(bullets) > 6:
        raise SkillError("invalid_bullets", "summary.bullets must have at most six items.")
    add_text(slide, name="Presentation Title [summary]", x=.85, y=.8, width=11.5, height=1.4,
             text=title, size=28 if len(title)<55 else 23, color=INK, bold=True)
    height = min(1.1, 4.15 / max(1, len(bullets)))
    for index, value in enumerate(bullets):
        text = required_text(value, label="summary.bullets", maximum=300)
        font_size = 22 if len(text) < 60 else 18 if len(text) < 120 else 14
        add_text(slide, name=f"Presentation Text {index} [summary]", x=1.05, y=2.45+index*height,
                 width=11.1, height=height-.08, text="•  "+text, size=font_size, color=INK)


def render_bullets(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "kicker", "bullets", "note"}, label="bullets slide")
    title = required_text(spec.get("title"), label="bullets.title", maximum=30)
    kicker = optional_text(spec.get("kicker"), label="bullets.kicker", maximum=30)
    bullets = parse_bullets(
        spec.get("bullets"),
        label="bullets.bullets",
        maximum_items=6,
        maximum_characters=60,
        maximum_total_characters=210,
    )
    note = optional_text(spec.get("note"), label="bullets.note", maximum=100)
    if profile == "academic-editorial":
        add_text(
            slide,
            name="ECNU Text Kicker [bullets]",
            x=0.95,
            y=1.2,
            width=2.3,
            height=0.28,
            text=(kicker or "EDITORIAL NOTE").upper(),
            size=8.5,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [bullets]",
            x=0.95,
            y=1.62,
            width=2.45,
            height=2.1,
            text=title,
            size=30,
            color=INK,
            bold=True,
        )
        add_rect(
            slide,
            name="ECNU Editorial Content Divider",
            x=3.62,
            y=1.18,
            width=0.035,
            height=5.34,
            fill=ECNU_RED,
        )
        add_bullet_list(
            slide,
            name="ECNU Text Bullets [bullets]",
            x=4.15,
            y=1.48,
            width=5.75,
            height=4.85,
            bullets=bullets,
            primary_size=19,
            secondary_size=16.5,
            item_spacing=14,
        )
        if note:
            add_text(
                slide,
                name="ECNU Text Note [bullets]",
                x=0.96,
                y=5.8,
                width=2.35,
                height=0.62,
                text=note,
                size=8.5,
                color=MUTED,
            )
        return
    if profile == "digital-campus":
        if kicker:
            add_text(
                slide,
                name="ECNU Text Kicker [bullets]",
                x=0.88,
                y=0.98,
                width=3.2,
                height=0.28,
                text=kicker.upper(),
                size=10,
                color=ECNU_RED,
                bold=True,
            )
        add_text(
            slide,
            name="ECNU Title [bullets]",
            x=0.88,
            y=1.3 if kicker else 1.1,
            width=8.05,
            height=0.72,
            text=title,
            size=32,
            color=INK,
            bold=True,
        )
        add_rule(slide, x=0.88, y=1.9 if kicker else 1.74, width=0.72)
        add_rect(
            slide,
            name="ECNU Digital Content Rail",
            x=0.96,
            y=2.17,
            width=0.14,
            height=4.15,
            fill=DIGITAL_PALE,
            line=DIGITAL_LINE,
        )
        add_rect(
            slide,
            name="ECNU Digital Content Accent",
            x=0.96,
            y=2.17,
            width=0.14,
            height=0.72,
            fill=ECNU_RED,
        )
        add_bullet_list(
            slide,
            name="ECNU Text Bullets [bullets]",
            x=1.38,
            y=2.22,
            width=8.45,
            height=3.92,
            bullets=bullets,
            normal_color=INK,
            accent_color=ECNU_RED,
            primary_size=22,
            secondary_size=17.5,
            item_spacing=14,
        )
        if note:
            add_text(
                slide,
                name="ECNU Text Note [bullets]",
                x=1.38,
                y=6.28,
                width=8.35,
                height=0.32,
                text=note,
                size=11,
                color=MUTED,
            )
        return
    add_title(slide, title, kicker=kicker, layout="bullets")
    add_bullet_list(slide, name="ECNU Text Bullets [bullets]", x=1.08, y=2.18, width=9.0, height=3.95, bullets=bullets)
    if note:
        add_text(
            slide,
            name="ECNU Text Note [bullets]",
            x=1.08,
            y=6.37,
            width=8.9,
            height=0.38,
            text=note,
            size=9,
            color=MUTED,
        )


def add_column_card(slide: Any, *, x: float, heading: str, bullets: list[dict[str, Any]], side: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(x), inches(2.12), inches(4.42), inches(4.15))
    card.name = f"ECNU Card {side} [two-column]"
    card.fill.solid()
    card.fill.fore_color.rgb = rgb(WHITE)
    card.line.color.rgb = rgb((214, 216, 219))
    card.line.width = points(0.8)
    add_rule(slide, x=x + 0.28, y=2.42, width=0.55)
    add_text(
        slide,
        name=f"ECNU Text Heading {side} [two-column]",
        x=x + 0.28,
        y=2.68,
        width=3.82,
        height=0.55,
        text=heading,
        size=19,
        color=INK,
        bold=True,
    )
    add_bullet_list(
        slide,
        name=f"ECNU Text Bullets {side} [two-column]",
        x=x + 0.3,
        y=3.37,
        width=3.72,
        height=2.52,
        bullets=bullets,
    )


def parse_column(value: Any, *, label: str) -> tuple[str, list[dict[str, Any]]]:
    if not isinstance(value, dict):
        raise SkillError("invalid_column", f"{label} must be an object.")
    reject_unknown(value, {"heading", "bullets"}, label=label)
    heading = required_text(value.get("heading"), label=f"{label}.heading", maximum=28)
    bullets = parse_bullets(
        value.get("bullets"),
        label=f"{label}.bullets",
        maximum_items=4,
        maximum_characters=42,
        maximum_total_characters=105,
    )
    return heading, bullets


def render_two_column(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "left", "right"}, label="two-column slide")
    title = required_text(spec.get("title"), label="two-column.title", maximum=30)
    left_heading, left_bullets = parse_column(spec.get("left"), label="two-column.left")
    right_heading, right_bullets = parse_column(spec.get("right"), label="two-column.right")
    if profile == "academic-editorial":
        add_title(slide, title, layout="two-column")
        add_text(
            slide,
            name="ECNU Text Heading Left [two-column]",
            x=1.0,
            y=2.25,
            width=3.85,
            height=0.48,
            text=f"01  {left_heading}",
            size=18,
            color=ECNU_RED,
            bold=True,
        )
        add_bullet_list(
            slide,
            name="ECNU Text Bullets Left [two-column]",
            x=1.0,
            y=2.98,
            width=3.9,
            height=3.25,
            bullets=left_bullets,
            primary_size=18,
            secondary_size=16,
        )
        add_rect(
            slide,
            name="ECNU Editorial Column Divider",
            x=5.1,
            y=2.22,
            width=0.035,
            height=4.1,
            fill=EDITORIAL_LINE,
        )
        add_text(
            slide,
            name="ECNU Text Heading Right [two-column]",
            x=5.55,
            y=2.25,
            width=3.85,
            height=0.48,
            text=f"02  {right_heading}",
            size=18,
            color=ECNU_RED,
            bold=True,
        )
        add_bullet_list(
            slide,
            name="ECNU Text Bullets Right [two-column]",
            x=5.55,
            y=2.98,
            width=3.9,
            height=3.25,
            bullets=right_bullets,
            primary_size=18,
            secondary_size=16,
        )
        return
    if profile == "digital-campus":
        add_title(slide, title, layout="two-column")
        for x, heading, bullets, side in (
            (0.92, left_heading, left_bullets, "Left"),
            (5.68, right_heading, right_bullets, "Right"),
        ):
            add_rect(
                slide,
                name=f"ECNU Digital Card {side} [two-column]",
                x=x,
                y=2.14,
                width=4.5,
                height=4.15,
                fill=DIGITAL_PALE,
                line=DIGITAL_LINE,
            )
            add_rect(
                slide,
                name=f"ECNU Digital Card Accent {side}",
                x=x,
                y=2.14,
                width=4.5,
                height=0.08,
                fill=ECNU_RED,
            )
            add_text(
                slide,
                name=f"ECNU Text Heading {side} [two-column]",
                x=x + 0.32,
                y=2.55,
                width=3.82,
                height=0.52,
                text=heading,
                size=21,
                color=DIGITAL_NAVY,
                bold=True,
            )
            add_bullet_list(
                slide,
                name=f"ECNU Text Bullets {side} [two-column]",
                x=x + 0.32,
                y=3.3,
                width=3.78,
                height=2.6,
                bullets=bullets,
                normal_color=INK,
                accent_color=ECNU_RED,
                primary_size=17.5,
                secondary_size=16,
            )
        return
    add_title(slide, title, layout="two-column")
    add_column_card(slide, x=0.95, heading=left_heading, bullets=left_bullets, side="Left")
    add_column_card(slide, x=5.65, heading=right_heading, bullets=right_bullets, side="Right")


def render_metrics(slide: Any, spec: dict[str, Any], profile: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    reject_unknown(spec, {"layout", "title", "metrics"}, label="metrics slide")
    title = required_text(spec.get("title"), label="metrics.title", maximum=30)
    metrics = spec.get("metrics")
    if not isinstance(metrics, list) or not 2 <= len(metrics) <= 4:
        raise SkillError("invalid_metrics", "metrics.metrics must contain 2 to 4 items.")
    parsed: list[dict[str, str]] = []
    for index, item in enumerate(metrics, start=1):
        if not isinstance(item, dict):
            raise SkillError("invalid_metric", f"metrics.metrics[{index}] must be an object.")
        reject_unknown(item, {"value", "label", "detail"}, label=f"metrics.metrics[{index}]")
        parsed.append(
            {
                "value": required_text(item.get("value"), label=f"metrics.metrics[{index}].value", maximum=8),
                "label": required_text(item.get("label"), label=f"metrics.metrics[{index}].label", maximum=20),
                "detail": optional_text(item.get("detail"), label=f"metrics.metrics[{index}].detail", maximum=50),
            }
        )
    if profile == "academic-editorial":
        add_title(slide, title, layout="metrics")
        row_height = 3.7 / len(parsed)
        for index, item in enumerate(parsed):
            y = 2.18 + index * row_height
            add_rule(slide, x=0.98, y=y, width=9.1, color=EDITORIAL_LINE)
            add_text(
                slide,
                name=f"ECNU Text Metric Value {index + 1} [metrics]",
                x=1.02,
                y=y + 0.22,
                width=1.65,
                height=0.65,
                text=item["value"],
                size=30,
                color=ECNU_RED,
                bold=True,
            )
            add_text(
                slide,
                name=f"ECNU Text Metric Label {index + 1} [metrics]",
                x=2.85,
                y=y + 0.24,
                width=2.5,
                height=0.5,
                text=item["label"],
                size=17,
                color=INK,
                bold=True,
            )
            if item["detail"]:
                add_text(
                    slide,
                    name=f"ECNU Text Metric Detail {index + 1} [metrics]",
                    x=5.6,
                    y=y + 0.27,
                    width=4.35,
                    height=0.48,
                    text=item["detail"],
                    size=12,
                    color=MUTED,
                    align=3,
                )
        return
    if profile == "digital-campus":
        add_title(slide, title, layout="metrics")
        gap = 0.18
        total_width = 9.46
        card_width = (total_width - gap * (len(parsed) - 1)) / len(parsed)
        for index, item in enumerate(parsed):
            x = 0.9 + index * (card_width + gap)
            add_rect(
                slide,
                name=f"ECNU Digital Metric Card {index + 1} [metrics]",
                x=x,
                y=2.25,
                width=card_width,
                height=3.82,
                fill=DIGITAL_PALE if index % 2 == 0 else WHITE,
                line=DIGITAL_LINE,
            )
            add_rect(
                slide,
                name=f"ECNU Digital Metric Accent {index + 1}",
                x=x,
                y=2.25,
                width=card_width,
                height=0.09,
                fill=ECNU_RED if index == 0 else DIGITAL_BLUE,
            )
            add_text(
                slide,
                name=f"ECNU Text Metric Value {index + 1} [metrics]",
                x=x + 0.24,
                y=2.8,
                width=card_width - 0.48,
                height=0.88,
                text=item["value"],
                size=33 if len(parsed) < 4 else 29,
                color=ECNU_RED if index == 0 else DIGITAL_NAVY,
                bold=True,
            )
            add_text(
                slide,
                name=f"ECNU Text Metric Label {index + 1} [metrics]",
                x=x + 0.24,
                y=3.93,
                width=card_width - 0.48,
                height=0.72,
                text=item["label"],
                size=17,
                color=INK,
                bold=True,
            )
            if item["detail"]:
                add_text(
                    slide,
                    name=f"ECNU Text Metric Detail {index + 1} [metrics]",
                    x=x + 0.24,
                    y=5.03,
                    width=card_width - 0.48,
                    height=0.62,
                    text=item["detail"],
                    size=12,
                    color=MUTED,
                )
        return
    add_title(slide, title, layout="metrics")
    gap = 0.22
    total_width = 9.52
    card_width = (total_width - gap * (len(parsed) - 1)) / len(parsed)
    for index, item in enumerate(parsed):
        x = 0.95 + index * (card_width + gap)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            inches(x),
            inches(2.35),
            inches(card_width),
            inches(3.45),
        )
        card.name = f"ECNU Card Metric {index + 1} [metrics]"
        card.fill.solid()
        card.fill.fore_color.rgb = rgb(WHITE)
        card.line.color.rgb = rgb((214, 216, 219))
        add_text(
            slide,
            name=f"ECNU Text Metric Value {index + 1} [metrics]",
            x=x + 0.24,
            y=2.75,
            width=card_width - 0.48,
            height=0.88,
            text=item["value"],
            size=31 if len(parsed) < 4 else 27,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name=f"ECNU Text Metric Label {index + 1} [metrics]",
            x=x + 0.24,
            y=3.78,
            width=card_width - 0.48,
            height=0.76,
            text=item["label"],
            size=17,
            color=INK,
            bold=True,
        )
        if item["detail"]:
            add_text(
                slide,
                name=f"ECNU Text Metric Detail {index + 1} [metrics]",
                x=x + 0.24,
                y=4.72,
                width=card_width - 0.48,
                height=0.7,
                text=item["detail"],
                size=12,
                color=MUTED,
            )


def parse_story_items(
    value: Any,
    *,
    label: str,
    minimum_items: int,
    maximum_items: int,
    key_name: str,
) -> list[dict[str, str]]:
    if not isinstance(value, list) or not minimum_items <= len(value) <= maximum_items:
        raise SkillError(
            f"invalid_{key_name}",
            f"{label} must contain {minimum_items} to {maximum_items} items.",
        )
    result: list[dict[str, str]] = []
    for index, item in enumerate(value, start=1):
        if not isinstance(item, dict):
            raise SkillError(f"invalid_{key_name}_item", f"{label}[{index}] must be an object.")
        if key_name == "timeline":
            reject_unknown(item, {"period", "title", "detail"}, label=f"{label}[{index}]")
            result.append(
                {
                    "marker": required_text(item.get("period"), label=f"{label}[{index}].period", maximum=16),
                    "title": required_text(item.get("title"), label=f"{label}[{index}].title", maximum=28),
                    "detail": optional_text(item.get("detail"), label=f"{label}[{index}].detail", maximum=52),
                }
            )
        elif key_name == "roadmap":
            reject_unknown(item, {"stage", "title", "detail"}, label=f"{label}[{index}]")
            result.append(
                {
                    "marker": required_text(item.get("stage"), label=f"{label}[{index}].stage", maximum=12),
                    "title": required_text(item.get("title"), label=f"{label}[{index}].title", maximum=24),
                    "detail": optional_text(item.get("detail"), label=f"{label}[{index}].detail", maximum=44),
                }
            )
        else:
            reject_unknown(item, {"title", "detail"}, label=f"{label}[{index}]")
            result.append(
                {
                    "marker": f"{index:02d}",
                    "title": required_text(item.get("title"), label=f"{label}[{index}].title", maximum=24),
                    "detail": optional_text(item.get("detail"), label=f"{label}[{index}].detail", maximum=52),
                }
            )
    return result


def render_timeline(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "events"}, label="timeline slide")
    title = required_text(spec.get("title"), label="timeline.title", maximum=30)
    events = parse_story_items(
        spec.get("events"),
        label="timeline.events",
        minimum_items=3,
        maximum_items=5,
        key_name="timeline",
    )
    add_title(slide, title, layout="timeline")
    if profile == "academic-editorial":
        row_height = 4.25 / len(events)
        add_rect(
            slide,
            name="ECNU Editorial Timeline Rail",
            x=2.48,
            y=2.12,
            width=0.035,
            height=4.28,
            fill=EDITORIAL_LINE,
        )
        for index, item in enumerate(events):
            y = 2.12 + index * row_height
            add_text(
                slide,
                name=f"ECNU Text Timeline Period {index + 1} [timeline]",
                x=0.98,
                y=y + 0.08,
                width=1.18,
                height=0.42,
                text=item["marker"],
                size=15,
                color=ECNU_RED,
                bold=True,
            )
            add_circle(
                slide,
                name=f"ECNU Timeline Node {index + 1} [timeline]",
                x=2.37,
                y=y + 0.18,
                diameter=0.25,
                fill=ECNU_RED if index == 0 else WHITE,
                line=ECNU_RED,
            )
            add_text(
                slide,
                name=f"ECNU Text Timeline Title {index + 1} [timeline]",
                x=2.88,
                y=y,
                width=2.9,
                height=0.45,
                text=item["title"],
                size=18,
                color=INK,
                bold=True,
            )
            if item["detail"]:
                add_text(
                    slide,
                    name=f"ECNU Text Timeline Detail {index + 1} [timeline]",
                    x=5.9,
                    y=y + 0.03,
                    width=4.05,
                    height=0.58,
                    text=item["detail"],
                    size=13,
                    color=MUTED,
                )
        return

    if profile == "digital-campus":
        add_rect(
            slide,
            name="ECNU Digital Timeline Band",
            x=0.98,
            y=3.05,
            width=9.1,
            height=0.16,
            fill=DIGITAL_PALE,
            line=DIGITAL_LINE,
        )
        node_fill = DIGITAL_NAVY
        node_diameter = 0.4
        node_y = 2.88
    else:
        add_rule(slide, x=1.02, y=3.1, width=9.0, color=EDITORIAL_LINE)
        node_fill = ECNU_RED
        node_diameter = 0.32
        node_y = 2.93
    column_width = 8.8 / len(events)
    for index, item in enumerate(events):
        center = 1.16 + column_width * (index + 0.5)
        add_circle(
            slide,
            name=f"ECNU Timeline Node {index + 1} [timeline]",
            x=center - node_diameter / 2,
            y=node_y,
            diameter=node_diameter,
            fill=ECNU_RED if index == 0 else node_fill,
        )
        add_text(
            slide,
            name=f"ECNU Text Timeline Period {index + 1} [timeline]",
            x=center - column_width / 2 + 0.08,
            y=2.25,
            width=column_width - 0.16,
            height=0.42,
            text=item["marker"],
            size=14,
            color=ECNU_RED,
            bold=True,
            align=1,
        )
        add_text(
            slide,
            name=f"ECNU Text Timeline Title {index + 1} [timeline]",
            x=center - column_width / 2 + 0.08,
            y=3.58,
            width=column_width - 0.16,
            height=0.78,
            text=item["title"],
            size=17,
            color=INK,
            bold=True,
            align=1,
        )
        if item["detail"]:
            add_text(
                slide,
                name=f"ECNU Text Timeline Detail {index + 1} [timeline]",
                x=center - column_width / 2 + 0.08,
                y=4.55,
                width=column_width - 0.16,
                height=1.02,
                text=item["detail"],
                size=12,
                color=MUTED,
                align=1,
            )


def render_feature_grid(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "items"}, label="feature-grid slide")
    title = required_text(spec.get("title"), label="feature-grid.title", maximum=30)
    items = parse_story_items(
        spec.get("items"),
        label="feature-grid.items",
        minimum_items=3,
        maximum_items=6,
        key_name="feature-grid",
    )
    add_title(slide, title, layout="feature-grid")
    columns = 3 if len(items) in {3, 5, 6} else 2
    rows = (len(items) + columns - 1) // columns
    cell_width = 9.08 / columns
    cell_height = 4.2 / rows
    for index, item in enumerate(items):
        column = index % columns
        row = index // columns
        x = 0.98 + column * cell_width
        y = 2.18 + row * cell_height
        if profile == "academic-editorial":
            add_rule(slide, x=x, y=y, width=cell_width - 0.24, color=EDITORIAL_LINE)
            text_offset = 0.12
        elif profile == "digital-campus":
            add_rect(
                slide,
                name=f"ECNU Digital Feature Field {index + 1} [feature-grid]",
                x=x,
                y=y,
                width=cell_width - 0.22,
                height=cell_height - 0.2,
                fill=DIGITAL_PALE if (row + column) % 2 == 0 else WHITE,
                line=DIGITAL_LINE,
            )
            text_offset = 0.3
        else:
            add_rect(
                slide,
                name=f"ECNU Feature Divider {index + 1} [feature-grid]",
                x=x,
                y=y + 0.02,
                width=0.06,
                height=cell_height - 0.34,
                fill=ECNU_RED if index == 0 else EDITORIAL_LINE,
            )
            text_offset = 0.22
        add_text(
            slide,
            name=f"ECNU Text Feature Number {index + 1} [feature-grid]",
            x=x + text_offset,
            y=y + 0.2,
            width=0.48,
            height=0.36,
            text=item["marker"],
            size=13,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name=f"ECNU Text Feature Title {index + 1} [feature-grid]",
            x=x + text_offset,
            y=y + 0.62,
            width=cell_width - text_offset - 0.43,
            height=0.5,
            text=item["title"],
            size=19,
            color=INK,
            bold=True,
        )
        if item["detail"]:
            add_text(
                slide,
                name=f"ECNU Text Feature Detail {index + 1} [feature-grid]",
                x=x + text_offset,
                y=y + 1.22,
                width=cell_width - text_offset - 0.43,
                height=max(0.62, cell_height - 1.55),
                text=item["detail"],
                size=13,
                color=MUTED,
            )


def render_roadmap(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "stages"}, label="roadmap slide")
    title = required_text(spec.get("title"), label="roadmap.title", maximum=30)
    stages = parse_story_items(
        spec.get("stages"),
        label="roadmap.stages",
        minimum_items=3,
        maximum_items=5,
        key_name="roadmap",
    )
    title_unit_limit = {3: 24.0, 4: 18.0, 5: 14.0}[len(stages)]
    for index, item in enumerate(stages, start=1):
        title_units = visual_text_units(item["title"])
        if title_units > title_unit_limit:
            raise SkillError(
                "roadmap_title_too_long",
                f"roadmap.stages[{index}].title is too long for a {len(stages)}-stage fixed layout.",
                details={
                    "stage_count": len(stages),
                    "display_units": round(title_units, 1),
                    "maximum_display_units": title_unit_limit,
                },
            )
    add_title(slide, title, layout="roadmap")
    if profile == "academic-editorial":
        row_height = 4.2 / len(stages)
        for index, item in enumerate(stages):
            y = 2.18 + index * row_height
            add_rect(
                slide,
                name=f"ECNU Roadmap Stage {index + 1} [roadmap]",
                x=0.98,
                y=y,
                width=1.35,
                height=row_height - 0.14,
                fill=ECNU_RED if index == 0 else EDITORIAL_WARM,
                line=ECNU_RED if index else None,
            )
            add_text(
                slide,
                name=f"ECNU Text Roadmap Marker {index + 1} [roadmap]",
                x=1.18,
                y=y + 0.22,
                width=0.95,
                height=0.36,
                text=item["marker"],
                size=14,
                color=WHITE if index == 0 else ECNU_RED,
                bold=True,
            )
            add_text(
                slide,
                name=f"ECNU Text Roadmap Title {index + 1} [roadmap]",
                x=2.75,
                y=y + 0.08,
                width=2.65,
                height=0.42,
                text=item["title"],
                size=18,
                color=INK,
                bold=True,
            )
            if item["detail"]:
                add_text(
                    slide,
                    name=f"ECNU Text Roadmap Detail {index + 1} [roadmap]",
                    x=5.55,
                    y=y + 0.1,
                    width=4.35,
                    height=0.52,
                    text=item["detail"],
                    size=13,
                    color=MUTED,
                )
        return

    gap = 0.16
    stage_width = (9.12 - gap * (len(stages) - 1)) / len(stages)
    for index, item in enumerate(stages):
        x = 0.98 + index * (stage_width + gap)
        rise = index * 0.22 if profile == "digital-campus" else 0
        y = 2.45 + (0.64 - min(rise, 0.64) if profile == "digital-campus" else 0)
        height = 3.55 + min(rise, 0.64) if profile == "digital-campus" else 3.9
        add_rect(
            slide,
            name=f"ECNU Roadmap Stage {index + 1} [roadmap]",
            x=x,
            y=y,
            width=stage_width,
            height=height,
            fill=DIGITAL_PALE if profile == "digital-campus" and index % 2 == 0 else WHITE,
            line=DIGITAL_LINE if profile == "digital-campus" else EDITORIAL_LINE,
        )
        add_rect(
            slide,
            name=f"ECNU Roadmap Accent {index + 1}",
            x=x,
            y=y,
            width=stage_width,
            height=0.08,
            fill=ECNU_RED if index == 0 else (DIGITAL_NAVY if profile == "digital-campus" else EDITORIAL_LINE),
        )
        add_text(
            slide,
            name=f"ECNU Text Roadmap Marker {index + 1} [roadmap]",
            x=x + 0.2,
            y=y + 0.38,
            width=stage_width - 0.4,
            height=0.36,
            text=item["marker"],
            size=13,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name=f"ECNU Text Roadmap Title {index + 1} [roadmap]",
            x=x + 0.2,
            y=y + 0.92,
            width=stage_width - 0.4,
            height=0.78,
            text=item["title"],
            size=18,
            color=INK,
            bold=True,
        )
        if item["detail"]:
            add_text(
                slide,
                name=f"ECNU Text Roadmap Detail {index + 1} [roadmap]",
                x=x + 0.2,
                y=y + 1.95,
                width=stage_width - 0.4,
                height=1.2,
                text=item["detail"],
                size=12,
                color=MUTED,
            )


def image_stream_and_size(path: Path) -> tuple[Any, int, int]:
    try:
        from PIL import Image

        image = Image.open(path)
        width, height = image.size
        if width <= 0 or height <= 0:
            raise ValueError("image dimensions must be positive")
        if path.suffix.lower() == ".webp":
            stream = io.BytesIO()
            converted = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            converted.save(stream, format="PNG")
            stream.seek(0)
            converted.close()
            image.close()
            return stream, width, height
        image.close()
        return str(path), width, height
    except SkillError:
        raise
    except Exception as exc:
        raise SkillError(
            "image_read_failed",
            "The slide image could not be read.",
            details={"path": str(path), "reason": str(exc)},
        ) from exc


def add_picture_contained(
    slide: Any,
    *,
    path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
) -> Any:
    source, pixel_width, pixel_height = image_stream_and_size(path)
    source_ratio = pixel_width / pixel_height
    box_ratio = width / height
    if source_ratio >= box_ratio:
        draw_width = width
        draw_height = width / source_ratio
    else:
        draw_height = height
        draw_width = height * source_ratio
    draw_x = x + (width - draw_width) / 2
    draw_y = y + (height - draw_height) / 2
    shape = slide.shapes.add_picture(
        source,
        inches(draw_x),
        inches(draw_y),
        width=inches(draw_width),
        height=inches(draw_height),
    )
    shape.name = "ECNU Content Image [image]"
    return shape


def render_image(slide: Any, spec: dict[str, Any], root: Path, profile: str) -> None:
    reject_unknown(spec, {"layout", "title", "image", "caption", "source"}, label="image slide")
    title = required_text(spec.get("title"), label="image.title", maximum=30)
    raw_image = spec.get("image")
    if not isinstance(raw_image, str):
        raise SkillError("invalid_image", "image.image must be an absolute project image path.")
    image_path = project_path(
        raw_image,
        root,
        label="image.image",
        must_exist=True,
        suffixes=(".png", ".jpg", ".jpeg", ".webp"),
    )
    validate_image_file(image_path)
    caption = optional_text(spec.get("caption"), label="image.caption", maximum=100)
    source = optional_text(spec.get("source"), label="image.source", maximum=100)
    if profile == "academic-editorial":
        add_title(slide, title, layout="image")
        add_picture_contained(slide, path=image_path, x=0.95, y=2.05, width=6.45, height=4.55)
        add_rect(
            slide,
            name="ECNU Editorial Image Divider",
            x=7.7,
            y=2.06,
            width=0.04,
            height=4.52,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Text Caption [image]",
            x=8.05,
            y=2.25,
            width=2.0,
            height=1.8,
            text=caption or "图像证据",
            size=15,
            color=INK,
            bold=True,
        )
        if source:
            add_text(
                slide,
                name="ECNU Text Source [image]",
                x=8.05,
                y=5.72,
                width=2.0,
                height=0.58,
                text=f"来源\n{source}",
                size=8,
                color=MUTED,
            )
        return
    if profile == "digital-campus":
        add_rect(
            slide,
            name="ECNU Digital Image Story Panel",
            x=0.86,
            y=1.45,
            width=3.1,
            height=5.28,
            fill=DIGITAL_PALE,
            line=DIGITAL_LINE,
        )
        add_rect(
            slide,
            name="ECNU Digital Image Story Accent",
            x=0.86,
            y=1.45,
            width=0.1,
            height=5.28,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Title [image]",
            x=1.28,
            y=1.9,
            width=2.25,
            height=1.25,
            text=title,
            size=27,
            color=INK,
            bold=True,
        )
        if caption:
            add_text(
                slide,
                name="ECNU Text Caption [image]",
                x=1.3,
                y=3.52,
                width=2.2,
                height=1.4,
                text=caption,
                size=15,
                color=MUTED,
            )
        if source:
            add_text(
                slide,
                name="ECNU Text Source [image]",
                x=1.3,
                y=5.85,
                width=2.2,
                height=0.38,
                text=f"来源：{source}",
                size=10,
                color=MUTED,
            )
        add_rect(
            slide,
            name="ECNU Digital Image Frame",
            x=4.18,
            y=1.45,
            width=6.12,
            height=5.28,
            fill=WHITE,
            line=DIGITAL_LINE,
        )
        add_picture_contained(slide, path=image_path, x=4.3, y=1.57, width=5.88, height=5.04)
        return
    add_title(slide, title, layout="image")
    add_picture_contained(slide, path=image_path, x=1.0, y=2.04, width=9.1, height=4.25)
    if caption:
        add_text(
            slide,
            name="ECNU Text Caption [image]",
            x=1.0,
            y=6.38,
            width=7.25,
            height=0.32,
            text=caption,
            size=9.5,
            color=INK,
        )
    if source:
        add_text(
            slide,
            name="ECNU Text Source [image]",
            x=8.0,
            y=6.4,
            width=2.1,
            height=0.28,
            text=f"来源：{source}",
            size=7.5,
            color=MUTED,
            align=3,
        )


def render_quote(slide: Any, spec: dict[str, Any], profile: str) -> None:
    reject_unknown(spec, {"layout", "quote", "attribution"}, label="quote slide")
    quote = required_text(spec.get("quote"), label="quote.quote", maximum=120)
    attribution = optional_text(spec.get("attribution"), label="quote.attribution", maximum=80)
    if profile == "academic-editorial":
        add_rect(
            slide,
            name="ECNU Editorial Quote Field",
            x=0.92,
            y=1.42,
            width=1.82,
            height=4.95,
            fill=EDITORIAL_WARM,
            line=EDITORIAL_LINE,
        )
        add_text(
            slide,
            name="ECNU Text Quote Mark [quote]",
            x=1.22,
            y=1.88,
            width=1.05,
            height=1.0,
            text="“",
            size=54,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Text Editorial Quote Label [quote]",
            x=1.25,
            y=5.35,
            width=1.15,
            height=0.34,
            text="KEY IDEA",
            size=8.5,
            color=MUTED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [quote]",
            x=3.28,
            y=1.9,
            width=6.55,
            height=2.75,
            text=quote,
            size=26,
            color=INK,
            bold=True,
        )
        if attribution:
            add_rule(slide, x=3.3, y=5.15, width=0.72)
            add_text(
                slide,
                name="ECNU Text Attribution [quote]",
                x=4.25,
                y=5.02,
                width=5.45,
                height=0.52,
                text=attribution,
                size=11,
                color=MUTED,
                align=3,
            )
        return
    if profile == "digital-campus":
        add_rect(
            slide,
            name="ECNU Digital Quote Field",
            x=0.92,
            y=1.48,
            width=1.72,
            height=4.92,
            fill=DIGITAL_PALE,
            line=DIGITAL_LINE,
        )
        add_text(
            slide,
            name="ECNU Text Quote Mark [quote]",
            x=1.25,
            y=1.83,
            width=0.9,
            height=1.0,
            text="“",
            size=54,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [quote]",
            x=3.15,
            y=2.14,
            width=6.65,
            height=2.25,
            text=quote,
            size=31,
            color=INK,
            bold=True,
        )
        if attribution:
            add_rect(
                slide,
                name="ECNU Digital Quote Signal",
                x=3.18,
                y=5.15,
                width=0.5,
                height=0.07,
                fill=ECNU_RED,
            )
            add_text(
                slide,
                name="ECNU Text Attribution [quote]",
                x=3.9,
                y=5.02,
                width=5.8,
                height=0.42,
                text=attribution,
                size=13,
                color=MUTED,
            )
        return
    add_text(
        slide,
        name="ECNU Text Quote Mark [quote]",
        x=0.92,
        y=1.55,
        width=1.0,
        height=0.9,
        text="“",
        size=52,
        color=ECNU_RED,
        bold=True,
    )
    add_text(
        slide,
        name="ECNU Title [quote]",
        x=1.45,
        y=2.12,
        width=8.25,
        height=2.5,
        text=quote,
        size=27,
        color=INK,
        bold=True,
    )
    if attribution:
        add_rule(slide, x=1.48, y=5.05, width=0.55)
        add_text(
            slide,
            name="ECNU Text Attribution [quote]",
            x=2.2,
            y=4.94,
            width=7.3,
            height=0.42,
            text=attribution,
            size=12,
            color=MUTED,
        )


def render_closing(slide: Any, spec: dict[str, Any], profile: str, theme: str) -> None:
    reject_unknown(spec, {"layout", "title", "subtitle", "contact"}, label="closing slide")
    title = required_text(spec.get("title"), label="closing.title", maximum=28)
    subtitle = optional_text(spec.get("subtitle"), label="closing.subtitle", maximum=80)
    contact = optional_text(spec.get("contact"), label="closing.contact", maximum=100)
    if profile == "academic-editorial":
        add_rect(
            slide,
            name="ECNU Editorial Closing Field",
            x=0.9,
            y=1.42,
            width=1.5,
            height=4.95,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Text Editorial Closing Label [closing]",
            x=1.17,
            y=1.88,
            width=0.95,
            height=0.55,
            text="LEARNING JOURNEY" if theme == "warm-education" else "EPILOGUE",
            size=8.5,
            color=WHITE,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [closing]",
            x=2.95,
            y=2.05,
            width=6.85,
            height=1.25,
            text=title,
            size=33,
            color=INK,
            bold=True,
        )
        add_rule(slide, x=2.98, y=3.55, width=1.65)
        if subtitle:
            add_text(
                slide,
                name="ECNU Text Subtitle [closing]",
                x=3.0,
                y=3.88,
                width=6.4,
                height=0.65,
                text=subtitle,
                size=16,
                color=MUTED,
            )
        if contact:
            add_text(
                slide,
                name="ECNU Text Contact [closing]",
                x=3.0,
                y=5.65,
                width=6.5,
                height=0.38,
                text=contact,
                size=10,
                color=ECNU_RED,
                bold=True,
            )
        return
    if profile == "digital-campus":
        add_rect(
            slide,
            name="ECNU Digital Closing Rail",
            x=0.88,
            y=1.48,
            width=0.14,
            height=4.9,
            fill=ECNU_RED,
        )
        add_text(
            slide,
            name="ECNU Text Digital Closing Label [closing]",
            x=1.38,
            y=1.92,
            width=2.8,
            height=0.3,
            text="NEXT · DIGITAL FUTURE" if theme == "digital-tech" else "NEXT · DIGITAL CAMPUS",
            size=10,
            color=ECNU_RED,
            bold=True,
        )
        add_text(
            slide,
            name="ECNU Title [closing]",
            x=1.34,
            y=2.42,
            width=7.6,
            height=1.08,
            text=title,
            size=40,
            color=INK,
            bold=True,
        )
        if subtitle:
            add_text(
                slide,
                name="ECNU Text Subtitle [closing]",
                x=1.38,
                y=3.84,
                width=7.1,
                height=0.62,
                text=subtitle,
                size=20,
                color=MUTED,
            )
        if contact:
            add_text(
                slide,
                name="ECNU Text Contact [closing]",
                x=1.38,
                y=5.45,
                width=6.5,
                height=0.42,
                text=contact,
                size=12,
                color=ECNU_RED,
                bold=True,
            )
        for index in range(3):
            add_rect(
                slide,
                name=f"ECNU Digital Closing Signal {index + 1}",
                x=8.45 + index * 0.36,
                y=5.35,
                width=0.22,
                height=0.22,
                fill=ECNU_RED if index == 0 else (DIGITAL_NAVY if index == 1 else DIGITAL_PALE),
                line=None if index < 2 else DIGITAL_LINE,
            )
        return
    add_rule(slide, x=0.95, y=2.25, width=0.95)
    add_text(
        slide,
        name="ECNU Title [closing]",
        x=0.95,
        y=2.62,
        width=8.8,
        height=1.1,
        text=title,
        size=34,
        color=INK,
        bold=True,
    )
    if subtitle:
        add_text(
            slide,
            name="ECNU Text Subtitle [closing]",
            x=0.98,
            y=3.92,
            width=8.5,
            height=0.55,
            text=subtitle,
            size=17,
            color=MUTED,
        )
    if contact:
        add_text(
            slide,
            name="ECNU Text Contact [closing]",
            x=0.98,
            y=5.55,
            width=8.5,
            height=0.42,
            text=contact,
            size=11,
            color=ECNU_RED,
            bold=True,
        )


def review_deck_plan(slides: list[Any]) -> list[dict[str, Any]]:
    layouts: list[str] = []
    for index, item in enumerate(slides, start=1):
        if not isinstance(item, dict):
            raise SkillError("invalid_slide", f"slides[{index}] must be an object.")
        layout = item.get("layout")
        if layout not in SUPPORTED_LAYOUTS:
            raise SkillError(
                "unsupported_layout",
                f"slides[{index}].layout is not supported.",
                details={"layout": layout, "supported": sorted(SUPPORTED_LAYOUTS)},
            )
        layouts.append(layout)

    slide_count = len(layouts)
    if slide_count < 12:
        return []

    counts = Counter(layouts)
    content_count = slide_count - sum(counts[layout] for layout in STRUCTURAL_LAYOUTS)
    quality_issues: list[dict[str, Any]] = []
    visual_count = sum(counts[layout] for layout in VISUAL_COMPOSITION_LAYOUTS)
    required_visuals = max(2, slide_count // 8)
    if visual_count < required_visuals:
        quality_issues.append(
            {
                "code": "too_few_visual_compositions",
                "message": "Long decks need timeline, roadmap, feature-grid, or image evidence instead of repeated text pages.",
                "count": visual_count,
                "minimum": required_visuals,
            }
        )

    section_limit = max(3, math.ceil(slide_count / 5))
    if counts["section"] > section_limit:
        quality_issues.append(
            {
                "code": "too_many_section_slides",
                "message": "Too many low-information section dividers interrupt the narrative.",
                "count": counts["section"],
                "maximum": section_limit,
            }
        )

    if content_count >= 8 and counts["bullets"] / content_count > 0.4:
        quality_issues.append(
            {
                "code": "too_many_bullet_slides",
                "message": "Bullet slides exceed 40% of content slides; use semantic visual layouts.",
                "count": counts["bullets"],
                "content_slide_count": content_count,
            }
        )

    if slide_count >= 16:
        distinct_content = {layout for layout in layouts if layout not in STRUCTURAL_LAYOUTS}
        if len(distinct_content) < 5:
            quality_issues.append(
                {
                    "code": "insufficient_layout_variety",
                    "message": "Long decks need at least five distinct content layouts.",
                    "count": len(distinct_content),
                    "minimum": 5,
                    "layouts": sorted(distinct_content),
                }
            )

    longest_run = 1
    run_layout = ""
    run_length = 0
    for layout in layouts:
        if layout == run_layout:
            run_length += 1
        else:
            run_layout = layout
            run_length = 1
        longest_run = max(longest_run, run_length)
    if longest_run > 2:
        quality_issues.append(
            {
                "code": "repeated_layout_run",
                "message": "The same layout may not appear more than twice in a row.",
                "longest_run": longest_run,
                "maximum": 2,
            }
        )

    return [
        {
            "severity": "warning",
            "category": "planning",
            **issue,
        }
        for issue in quality_issues
    ]


def presentation_metadata(presentation: Any, value: Any, profile: str, theme: str, branded: bool) -> None:
    if value is None:
        value = {}
    if not isinstance(value, dict):
        raise SkillError("invalid_metadata", "metadata must be an object.")
    reject_unknown(value, {"title", "author", "subject", "language"}, label="metadata")
    title = optional_text(value.get("title"), label="metadata.title", maximum=160)
    author = optional_text(value.get("author"), label="metadata.author", maximum=100)
    subject = optional_text(value.get("subject"), label="metadata.subject", maximum=160)
    language = optional_text(value.get("language"), label="metadata.language", maximum=20) or "zh-CN"
    presentation.core_properties.title = title or ("ECNU Presentation" if branded else "Presentation")
    presentation.core_properties.author = author or "DSH Artifact Services"
    presentation.core_properties.subject = subject
    presentation.core_properties.language = language
    parts = [PRESENTATION_MARKER, "structured", f"profile={profile}"]
    if theme:
        parts.append(f"theme={theme}")
    if branded:
        parts.append("ECNU")
    presentation.core_properties.keywords = "; ".join(parts)


def main() -> int:
    args = parse_args()
    root = project_root(args.project_root)
    output = project_path(args.output, root, label="--output", suffixes=(".pptx",))
    spec = read_json_spec(inline=args.spec_json, file_path=args.spec_file, root=root)
    reject_unknown(spec, {"theme", "profile", "metadata", "footer", "slides"}, label="presentation specification")
    theme, profile, branded = presentation_style(spec.get("theme"), spec.get("profile"))
    apply_theme_palette(theme)
    footer = optional_text(spec.get("footer"), label="footer", maximum=100)
    slides = spec.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 40:
        raise SkillError("invalid_slides", "slides must contain between 1 and 40 slide objects.")
    plan_warnings = review_deck_plan(slides)

    pptx = require_pptx()
    presentation = pptx.Presentation()
    presentation.slide_width = inches(SLIDE_WIDTH_IN)
    presentation.slide_height = inches(SLIDE_HEIGHT_IN)
    presentation_metadata(presentation, spec.get("metadata"), profile, theme, branded)
    background = asset_path("ecnu-background.jpg") if branded else None
    logo_a = asset_path("ecnu-logo-form-a.png") if branded else None
    logo_c = asset_path("ecnu-logo-form-c.png") if branded else None
    layout_counts: Counter[str] = Counter()

    renderers = {
        "summary": lambda slide, item: render_summary(slide, item, profile),
        "cover": lambda slide, item: render_cover(slide, item, profile, theme),
        "section": lambda slide, item: render_section(slide, item, profile),
        "bullets": lambda slide, item: render_bullets(slide, item, profile),
        "two-column": lambda slide, item: render_two_column(slide, item, profile),
        "metrics": lambda slide, item: render_metrics(slide, item, profile),
        "timeline": lambda slide, item: render_timeline(slide, item, profile),
        "feature-grid": lambda slide, item: render_feature_grid(slide, item, profile),
        "roadmap": lambda slide, item: render_roadmap(slide, item, profile),
        "image": lambda slide, item: render_image(slide, item, root, profile),
        "quote": lambda slide, item: render_quote(slide, item, profile),
        "closing": lambda slide, item: render_closing(slide, item, profile, theme),
    }

    for index, slide_spec in enumerate(slides, start=1):
        layout = slide_spec.get("layout")
        slide = add_base(
            presentation,
            profile=profile,
            theme=theme,
            branded=branded,
            layout=layout,
            page_number=index,
            footer=footer,
            background=background,
            logo_a=logo_a,
            logo_c=logo_c,
        )
        visible_spec = {key: value for key, value in slide_spec.items() if key != "speaker_notes"}
        renderers[layout](slide, visible_spec)
        notes = slide_spec.get("speaker_notes", "")
        if not isinstance(notes, str) or len(notes) > 100000:
            raise SkillError("invalid_speaker_notes", "speaker_notes must be text with at most 100000 characters.")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        layout_counts[layout] += 1

    atomic_save(presentation, output, overwrite=args.overwrite)
    emit(
        {
            "ok": True,
            "operation": "create",
            **path_result(output, root),
            "slide_count": len(slides),
            "layout_counts": dict(sorted(layout_counts.items())),
            "profile": profile,
            "theme": theme or "legacy-ecnu-profile",
            "supported_themes": sorted(SUPPORTED_THEMES),
            "default_theme": DEFAULT_THEME,
            "supported_profiles": sorted(SUPPORTED_PROFILES),
            "default_profile": DEFAULT_PROFILE,
            "brand": "ECNU fixed v3" if branded else "generic theme",
            "plan_quality_pass": not plan_warnings,
            "plan_warnings": plan_warnings,
            "requires_structural_validation": True,
            "requires_visual_review_for_final_delivery": True,
        }
    )
    return 0


if __name__ == "__main__":
    run("create", main)
