"""Shared validation, rendering, and JSON helpers for managed DOCX modules."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterable


EXIT_USAGE = 2
EXIT_NOT_FOUND = 3
EXIT_DOCUMENT = 4
EXIT_OUTPUT = 5
EXIT_VALIDATION = 6
EXIT_DEPENDENCY = 10

ALIGNMENTS = {"left", "center", "right", "justify"}
HEX_COLOR = re.compile(r"^[0-9A-Fa-f]{6}$")
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
MAX_DOCX_BYTES = 256 * 1024 * 1024
MAX_SPEC_BYTES = 4 * 1024 * 1024
MAX_IMAGE_BYTES = 32 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
MAX_ZIP_ENTRIES = 10_000
MAX_ZIP_ENTRY_BYTES = 128 * 1024 * 1024
MAX_ZIP_EXPANDED_BYTES = 512 * 1024 * 1024


class SkillError(Exception):
    def __init__(
        self,
        code: str,
        message: str,
        *,
        exit_code: int = EXIT_USAGE,
        details: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.exit_code = exit_code
        self.details = details or {}


class JsonArgumentParser(argparse.ArgumentParser):
    def error(self, message: str) -> None:
        raise SkillError("invalid_arguments", message, exit_code=EXIT_USAGE)


def emit(payload: dict[str, Any], *, stream: Any = sys.stdout) -> None:
    stream.write(json.dumps(payload, ensure_ascii=False) + "\n")
    stream.flush()


def run(operation: str, function: Callable[[], int | None]) -> None:
    try:
        result = function()
    except SkillError as exc:
        payload: dict[str, Any] = {
            "ok": False,
            "operation": operation,
            "error": {"code": exc.code, "message": exc.message},
        }
        if exc.details:
            payload["error"]["details"] = exc.details
        emit(payload, stream=sys.stderr)
        raise SystemExit(exc.exit_code)
    except KeyboardInterrupt:
        emit(
            {
                "ok": False,
                "operation": operation,
                "error": {"code": "cancelled", "message": "Operation was cancelled."},
            },
            stream=sys.stderr,
        )
        raise SystemExit(130)
    except Exception as exc:
        emit(
            {
                "ok": False,
                "operation": operation,
                "error": {"code": "internal_error", "message": str(exc)},
            },
            stream=sys.stderr,
        )
        raise SystemExit(1)
    raise SystemExit(0 if result is None else result)


def require_docx() -> Any:
    try:
        import docx
    except ImportError as exc:
        raise SkillError(
            "dependency_missing",
            "python-docx is required in the managed Python environment.",
            exit_code=EXIT_DEPENDENCY,
            details={"package": "python-docx", "recovery": "install python/requirements.txt into the interpreter configured by DSH_OFFICE_PYTHON"},
        ) from exc
    return docx


def project_root(raw: str) -> Path:
    path = Path(raw)
    if not path.is_absolute():
        raise SkillError("absolute_path_required", "--project-root must be an absolute path.")
    try:
        resolved = path.resolve(strict=True)
    except FileNotFoundError as exc:
        raise SkillError(
            "project_not_found",
            "The project root does not exist.",
            exit_code=EXIT_NOT_FOUND,
            details={"path": str(path)},
        ) from exc
    if not resolved.is_dir():
        raise SkillError("invalid_project_root", "The project root is not a directory.")
    return resolved


def project_path(
    raw: str,
    root: Path,
    *,
    label: str,
    must_exist: bool = False,
    suffix: str | None = None,
    suffixes: set[str] | None = None,
) -> Path:
    path = Path(raw)
    if not path.is_absolute():
        raise SkillError("absolute_path_required", f"{label} must be an absolute path.")
    resolved = path.resolve(strict=False)
    try:
        common = os.path.commonpath((str(root), str(resolved)))
    except ValueError as exc:
        raise SkillError("path_outside_project", f"{label} is outside the project root.") from exc
    if os.path.normcase(common) != os.path.normcase(str(root)):
        raise SkillError(
            "path_outside_project",
            f"{label} is outside the project root.",
            details={"path": str(resolved)},
        )
    allowed = {suffix.lower()} if suffix else {item.lower() for item in suffixes or set()}
    if allowed and resolved.suffix.lower() not in allowed:
        raise SkillError(
            "invalid_file_type",
            f"{label} must use one of these extensions: {', '.join(sorted(allowed))}.",
            details={"path": str(resolved)},
        )
    if must_exist:
        if not resolved.exists():
            raise SkillError(
                "file_not_found",
                f"{label} does not exist.",
                exit_code=EXIT_NOT_FOUND,
                details={"path": str(resolved)},
            )
        if not resolved.is_file():
            raise SkillError("not_a_file", f"{label} is not a file.")
        size = resolved.stat().st_size
        limit = None
        if resolved.suffix.lower() == ".docx":
            limit = MAX_DOCX_BYTES
        elif resolved.suffix.lower() == ".json":
            limit = MAX_SPEC_BYTES
        elif resolved.suffix.lower() in IMAGE_SUFFIXES:
            limit = MAX_IMAGE_BYTES
        if limit is not None and size > limit:
            raise SkillError(
                "file_too_large",
                f"{label} exceeds the managed file-size limit.",
                exit_code=EXIT_DOCUMENT,
                details={"path": str(resolved), "size_bytes": size, "max_bytes": limit},
            )
    return resolved


def read_json_spec(*, inline: str | None, file_path: str | None, root: Path) -> dict[str, Any]:
    if bool(inline) == bool(file_path):
        raise SkillError("invalid_arguments", "Provide exactly one of --spec-json or --spec-file.")
    source = "inline JSON"
    text = inline
    if file_path:
        spec_path = project_path(file_path, root, label="--spec-file", must_exist=True, suffix=".json")
        source = str(spec_path)
        try:
            text = spec_path.read_text(encoding="utf-8")
        except OSError as exc:
            raise SkillError(
                "spec_read_failed",
                "The JSON specification could not be read.",
                exit_code=EXIT_NOT_FOUND,
                details={"path": str(spec_path), "reason": str(exc)},
            ) from exc
    try:
        parsed = json.loads(text or "")
    except json.JSONDecodeError as exc:
        raise SkillError(
            "invalid_json",
            "The JSON specification is invalid.",
            details={"source": source, "line": exc.lineno, "column": exc.colno, "reason": exc.msg},
        ) from exc
    if not isinstance(parsed, dict):
        raise SkillError("invalid_spec", "The JSON specification must be an object.")
    return parsed


def non_empty_text(value: Any, *, label: str, maximum: int = 100_000) -> str:
    if not isinstance(value, str) or not value.strip():
        raise SkillError("invalid_text", f"{label} must be a non-empty string.")
    if len(value) > maximum:
        raise SkillError("text_too_long", f"{label} exceeds {maximum} characters.")
    return value


def optional_text(value: Any, *, label: str, maximum: int = 100_000) -> str:
    if not isinstance(value, str):
        raise SkillError("invalid_text", f"{label} must be a string.")
    if len(value) > maximum:
        raise SkillError("text_too_long", f"{label} exceeds {maximum} characters.")
    return value


def finite_number(value: Any, *, label: str, minimum: float, maximum: float) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise SkillError("invalid_number", f"{label} must be a number.")
    number = float(value)
    if not minimum <= number <= maximum:
        raise SkillError("invalid_number", f"{label} must be between {minimum} and {maximum}.")
    return number


def alignment(value: Any, *, label: str) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str) or value not in ALIGNMENTS:
        raise SkillError("invalid_alignment", f"{label} must be one of {sorted(ALIGNMENTS)}.")
    return value


def atomic_save(document: Any, output: Path, *, overwrite: bool) -> None:
    if output.exists() and not overwrite:
        raise SkillError(
            "output_exists",
            "The output file already exists; pass --overwrite only after confirming replacement.",
            exit_code=EXIT_OUTPUT,
            details={"path": str(output)},
        )
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise SkillError(
            "output_directory_failed",
            "The output directory could not be created.",
            exit_code=EXIT_OUTPUT,
            details={"path": str(output.parent), "reason": str(exc)},
        ) from exc
    descriptor = -1
    temporary = ""
    try:
        descriptor, temporary = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=".docx", dir=str(output.parent))
        os.close(descriptor)
        descriptor = -1
        document.save(temporary)
        if os.path.getsize(temporary) > MAX_DOCX_BYTES:
            raise SkillError(
                "file_too_large",
                "The generated DOCX exceeds the managed size limit.",
                exit_code=EXIT_OUTPUT,
                details={"size_bytes": os.path.getsize(temporary), "max_bytes": MAX_DOCX_BYTES},
            )
        os.replace(temporary, output)
    except SkillError:
        raise
    except Exception as exc:
        raise SkillError(
            "document_write_failed",
            "The document could not be written.",
            exit_code=EXIT_OUTPUT,
            details={"path": str(output), "reason": str(exc)},
        ) from exc
    finally:
        if descriptor >= 0:
            os.close(descriptor)
        if temporary and os.path.exists(temporary):
            try:
                os.remove(temporary)
            except OSError:
                pass


def path_result(path: Path, root: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "path": str(path),
        "project_relative_path": path.relative_to(root).as_posix(),
    }
    if path.exists():
        result["size_bytes"] = path.stat().st_size
    return result


def validate_ooxml_archive(path: Path) -> None:
    if not zipfile.is_zipfile(path):
        raise SkillError("invalid_docx_container", "The file is not a DOCX ZIP container.", exit_code=EXIT_DOCUMENT)
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise SkillError(
                    "archive_too_complex",
                    "The DOCX contains too many archive entries.",
                    exit_code=EXIT_DOCUMENT,
                    details={"entry_count": len(entries), "max_entries": MAX_ZIP_ENTRIES},
                )
            expanded = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise SkillError("encrypted_archive", "Encrypted DOCX parts are not supported.", exit_code=EXIT_DOCUMENT)
                if entry.file_size > MAX_ZIP_ENTRY_BYTES:
                    raise SkillError(
                        "archive_entry_too_large",
                        "A DOCX archive entry exceeds the managed expansion limit.",
                        exit_code=EXIT_DOCUMENT,
                        details={"entry": entry.filename, "size_bytes": entry.file_size, "max_bytes": MAX_ZIP_ENTRY_BYTES},
                    )
                expanded += entry.file_size
                if expanded > MAX_ZIP_EXPANDED_BYTES:
                    raise SkillError(
                        "archive_expansion_too_large",
                        "The DOCX expanded content exceeds the managed limit.",
                        exit_code=EXIT_DOCUMENT,
                        details={"expanded_bytes": expanded, "max_bytes": MAX_ZIP_EXPANDED_BYTES},
                    )
    except zipfile.BadZipFile as exc:
        raise SkillError("invalid_docx_container", "The DOCX ZIP container is invalid.", exit_code=EXIT_DOCUMENT) from exc


def validate_image_file(path: Path) -> None:
    try:
        from PIL import Image
    except ImportError as exc:
        raise SkillError(
            "dependency_missing",
            "Pillow is required to validate document images.",
            exit_code=EXIT_DEPENDENCY,
            details={"package": "Pillow"},
        ) from exc
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as image:
                width, height = image.size
                if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                    raise SkillError(
                        "image_too_large",
                        "The image exceeds the managed pixel limit.",
                        exit_code=EXIT_DOCUMENT,
                        details={"width": width, "height": height, "max_pixels": MAX_IMAGE_PIXELS},
                    )
                image.verify()
    except SkillError:
        raise
    except Exception as exc:
        raise SkillError(
            "image_read_failed",
            "The image could not be safely decoded.",
            exit_code=EXIT_DOCUMENT,
            details={"path": str(path), "reason": str(exc)},
        ) from exc


def iter_paragraphs(document: Any) -> Iterable[Any]:
    yield from document.paragraphs
    seen_cells: set[int] = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen_cells:
                    continue
                seen_cells.add(marker)
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def reject_unknown(mapping: dict[str, Any], allowed: set[str], *, label: str) -> None:
    unknown = sorted(set(mapping) - allowed)
    if unknown:
        raise SkillError(
            "unknown_fields",
            f"{label} contains unsupported fields.",
            details={"fields": unknown},
        )


def apply_alignment(paragraph: Any, value: str | None, docx: Any) -> None:
    if value is None:
        return
    mapping = {
        "left": docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT,
        "center": docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER,
        "right": docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    paragraph.alignment = mapping[value]


def apply_run(run: Any, spec: dict[str, Any], *, label: str, docx: Any) -> None:
    reject_unknown(
        spec,
        {"text", "bold", "italic", "underline", "font_size_pt", "font", "east_asia_font", "color"},
        label=label,
    )
    for field in ("bold", "italic", "underline"):
        if field in spec:
            value = spec[field]
            if not isinstance(value, bool):
                raise SkillError("invalid_run_format", f"{label}.{field} must be boolean.")
            setattr(run, field, value)
    if "font_size_pt" in spec:
        run.font.size = docx.shared.Pt(
            finite_number(spec["font_size_pt"], label=f"{label}.font_size_pt", minimum=6, maximum=96)
        )
    if "font" in spec:
        run.font.name = non_empty_text(spec["font"], label=f"{label}.font", maximum=100)
    if "east_asia_font" in spec:
        east_asia = non_empty_text(spec["east_asia_font"], label=f"{label}.east_asia_font", maximum=100)
        run._element.get_or_add_rPr().rFonts.set(docx.oxml.ns.qn("w:eastAsia"), east_asia)
    if "color" in spec:
        color = non_empty_text(spec["color"], label=f"{label}.color", maximum=6)
        if not HEX_COLOR.fullmatch(color):
            raise SkillError("invalid_color", f"{label}.color must be a six-digit RGB hex value.")
        run.font.color.rgb = docx.shared.RGBColor.from_string(color.upper())


def add_text(paragraph: Any, spec: dict[str, Any], *, label: str, docx: Any) -> None:
    has_text = "text" in spec
    has_runs = "runs" in spec
    if has_text == has_runs:
        raise SkillError("invalid_text_spec", f"{label} must contain exactly one of text or runs.")
    if has_text:
        run = paragraph.add_run(optional_text(spec["text"], label=f"{label}.text"))
        # List items commonly use the compact {text, bold, ...} form.  Keep
        # the same formatting vocabulary as rich runs instead of forcing a
        # model to wrap a one-run item in a runs array.
        compact = {key: spec[key] for key in (
            "text", "bold", "italic", "underline", "font_size_pt",
            "font", "east_asia_font", "color",
        ) if key in spec}
        apply_run(run, compact, label=label, docx=docx)
        return
    runs = spec["runs"]
    if not isinstance(runs, list) or not runs:
        raise SkillError("invalid_runs", f"{label}.runs must be a non-empty array.")
    for index, run_spec in enumerate(runs, start=1):
        if not isinstance(run_spec, dict):
            raise SkillError("invalid_run", f"{label}.runs[{index}] must be an object.")
        text = optional_text(run_spec.get("text"), label=f"{label}.runs[{index}].text")
        run = paragraph.add_run(text)
        apply_run(run, run_spec, label=f"{label}.runs[{index}]", docx=docx)


def set_cell_shading(cell: Any, fill: str, docx: Any) -> None:
    shading = docx.oxml.OxmlElement("w:shd")
    shading.set(docx.oxml.ns.qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_blocks(document: Any, blocks: Any, *, root: Path, docx: Any, label: str = "blocks") -> dict[str, int]:
    if not isinstance(blocks, list) or not blocks:
        raise SkillError("invalid_blocks", f"{label} must be a non-empty array.")
    if len(blocks) > 2_000:
        raise SkillError("too_many_blocks", f"{label} must contain at most 2000 blocks.")
    counts: dict[str, int] = {}
    for index, block in enumerate(blocks, start=1):
        block_label = f"{label}[{index}]"
        if not isinstance(block, dict):
            raise SkillError("invalid_block", f"{block_label} must be an object.")
        kind = block.get("type")
        if not isinstance(kind, str):
            raise SkillError("invalid_block", f"{block_label}.type must be a string.")

        if kind == "title":
            reject_unknown(block, {"type", "text", "alignment"}, label=block_label)
            paragraph = document.add_heading(non_empty_text(block.get("text"), label=f"{block_label}.text"), level=0)
            apply_alignment(paragraph, alignment(block.get("alignment", "center"), label=f"{block_label}.alignment"), docx)
        elif kind == "heading":
            reject_unknown(block, {"type", "level", "text", "alignment"}, label=block_label)
            level = block.get("level")
            if isinstance(level, bool) or not isinstance(level, int) or not 1 <= level <= 6:
                raise SkillError("invalid_heading_level", f"{block_label}.level must be an integer from 1 to 6.")
            paragraph = document.add_heading(non_empty_text(block.get("text"), label=f"{block_label}.text"), level=level)
            apply_alignment(paragraph, alignment(block.get("alignment"), label=f"{block_label}.alignment"), docx)
        elif kind == "paragraph":
            reject_unknown(block, {"type", "text", "runs", "alignment"}, label=block_label)
            paragraph = document.add_paragraph()
            add_text(paragraph, block, label=block_label, docx=docx)
            apply_alignment(paragraph, alignment(block.get("alignment"), label=f"{block_label}.alignment"), docx)
        elif kind == "list":
            reject_unknown(block, {"type", "ordered", "items"}, label=block_label)
            ordered = block.get("ordered", False)
            items = block.get("items")
            if not isinstance(ordered, bool):
                raise SkillError("invalid_list", f"{block_label}.ordered must be boolean.")
            if not isinstance(items, list) or not items:
                raise SkillError("invalid_list", f"{block_label}.items must be a non-empty array.")
            for item_index, item in enumerate(items, start=1):
                item_label = f"{block_label}.items[{item_index}]"
                level = 0
                if isinstance(item, str):
                    text_spec = {"text": optional_text(item, label=item_label)}
                elif isinstance(item, dict):
                    reject_unknown(item, {
                        "text", "runs", "level", "bold", "italic", "underline",
                        "font_size_pt", "font", "east_asia_font", "color",
                    }, label=item_label)
                    text_spec = {key: value for key, value in item.items() if key != "level"}
                    level = item.get("level", 0)
                    if isinstance(level, bool) or not isinstance(level, int) or not 0 <= level <= 2:
                        raise SkillError("invalid_list_level", f"{item_label}.level must be 0, 1, or 2.")
                else:
                    raise SkillError("invalid_list_item", f"{item_label} must be a string or object.")
                style = "List Number" if ordered else "List Bullet"
                if level:
                    style += f" {level + 1}"
                paragraph = document.add_paragraph(style=style)
                add_text(paragraph, text_spec, label=item_label, docx=docx)
        elif kind == "table":
            reject_unknown(block, {"type", "rows", "header", "column_widths_cm"}, label=block_label)
            rows = block.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(row, list) and row for row in rows):
                raise SkillError("invalid_table", f"{block_label}.rows must be a non-empty array of non-empty rows.")
            columns = len(rows[0])
            if columns > 50 or len(rows) > 5_000:
                raise SkillError("table_too_large", f"{block_label} exceeds 5000 rows or 50 columns.")
            if any(len(row) != columns for row in rows):
                raise SkillError("ragged_table", f"{block_label}.rows must all have the same number of cells.")
            table = document.add_table(rows=len(rows), cols=columns)
            table.style = "Table Grid"
            table.autofit = True
            header = block.get("header", False)
            if not isinstance(header, bool):
                raise SkillError("invalid_table", f"{block_label}.header must be boolean.")
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    if value is None:
                        text = ""
                    elif isinstance(value, (str, int, float, bool)):
                        text = str(value)
                    else:
                        raise SkillError(
                            "invalid_table_cell",
                            f"{block_label}.rows[{row_index + 1}][{column_index + 1}] must be scalar or null.",
                        )
                    cell = table.cell(row_index, column_index)
                    cell.text = text
                    if header and row_index == 0:
                        set_cell_shading(cell, "E9EEF5", docx)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            widths = block.get("column_widths_cm")
            if widths is not None:
                if not isinstance(widths, list) or len(widths) != columns:
                    raise SkillError("invalid_column_widths", f"{block_label}.column_widths_cm must match column count.")
                for column_index, width in enumerate(widths):
                    width_cm = finite_number(width, label=f"{block_label}.column_widths_cm[{column_index + 1}]", minimum=0.5, maximum=40)
                    for cell in table.columns[column_index].cells:
                        cell.width = docx.shared.Cm(width_cm)
        elif kind == "image":
            reject_unknown(block, {"type", "path", "width_cm", "caption"}, label=block_label)
            image_path = project_path(
                non_empty_text(block.get("path"), label=f"{block_label}.path"),
                root,
                label=f"{block_label}.path",
                must_exist=True,
                suffixes=IMAGE_SUFFIXES,
            )
            validate_image_file(image_path)
            width = block.get("width_cm")
            width_arg = None if width is None else docx.shared.Cm(
                finite_number(width, label=f"{block_label}.width_cm", minimum=0.5, maximum=40)
            )
            try:
                document.add_picture(str(image_path), width=width_arg)
            except Exception as exc:
                raise SkillError(
                    "image_read_failed",
                    "The image could not be added to the document.",
                    exit_code=EXIT_DOCUMENT,
                    details={"path": str(image_path), "reason": str(exc)},
                ) from exc
            caption = block.get("caption")
            if caption is not None:
                paragraph = document.add_paragraph(optional_text(caption, label=f"{block_label}.caption"))
                apply_alignment(paragraph, "center", docx)
        elif kind == "page_break":
            reject_unknown(block, {"type"}, label=block_label)
            document.add_page_break()
        else:
            raise SkillError(
                "unsupported_block",
                f"{block_label}.type is not supported.",
                details={"type": kind, "supported": ["title", "heading", "paragraph", "list", "table", "image", "page_break"]},
            )
        counts[kind] = counts.get(kind, 0) + 1
    return counts


def apply_properties(document: Any, value: Any, *, label: str = "properties") -> None:
    if not isinstance(value, dict):
        raise SkillError("invalid_properties", f"{label} must be an object.")
    aliases = {"author": "author", "creator": "author", "title": "title", "subject": "subject", "keywords": "keywords", "comments": "comments"}
    for key, content in value.items():
        if key not in aliases:
            raise SkillError("unsupported_property", f"{label}.{key} is not supported.")
        if content is not None and not isinstance(content, str):
            raise SkillError("invalid_property", f"{label}.{key} must be a string or null.")
        setattr(document.core_properties, aliases[key], content or "")
