from __future__ import annotations

import argparse
import base64
import hashlib
import html
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


MAX_SHEETS = 12
MAX_ROWS = 240
MAX_COLUMNS = 60
MAX_SLIDES = 80


def escaped(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def data_url(content_type: str, blob: bytes) -> str:
    return f"data:{content_type};base64,{base64.b64encode(blob).decode('ascii')}"


def color_value(color: Any) -> str | None:
    try:
        value = str(color.rgb or "")
    except (AttributeError, TypeError, ValueError):
        return None
    if len(value) == 8:
        value = value[2:]
    return f"#{value}" if len(value) == 6 else None


def document_shell(title: str, kind: str, body: str, extra_style: str = "") -> str:
    safe_title = escaped(title)
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; img-src data:; style-src 'unsafe-inline'; font-src data:">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{safe_title}</title><style>
:root{{color-scheme:light;--paper:#fff;--ink:#252525;--muted:#6d7480;--line:#d9dde5;--canvas:#eef0f4;--accent:#8b2232}}
*{{box-sizing:border-box}}html,body{{margin:0;min-height:100%;background:var(--canvas);color:var(--ink);font-family:"Microsoft YaHei","Segoe UI",sans-serif}}
body{{padding:18px}}.preview-meta{{max-width:1100px;margin:0 auto 12px;color:var(--muted);font-size:12px}}
.paper{{width:min(100%,860px);min-height:1080px;margin:0 auto 18px;padding:58px 68px;background:var(--paper);box-shadow:0 5px 24px #1f29371f}}
.paper p{{margin:.45em 0;line-height:1.72;white-space:pre-wrap}}.paper h1,.paper h2,.paper h3{{line-height:1.35;margin:1em 0 .5em}}
.paper img{{max-width:100%;height:auto}}table{{border-collapse:collapse}}.paper table{{width:100%;margin:14px 0}}th,td{{border:1px solid var(--line);padding:6px 8px;vertical-align:top}}
.page-break{{border:0;border-top:2px dashed var(--line);margin:34px -20px}}.note{{padding:9px 12px;border-radius:8px;background:#fff8e7;color:#705a1d;font-size:12px}}
.sheet{{width:max-content;min-width:min(100%,900px);max-width:none;margin:0 auto 22px;background:#fff;box-shadow:0 4px 18px #1f29371a}}
.sheet h2{{position:sticky;left:0;margin:0;padding:12px 14px;border-bottom:1px solid var(--line);font-size:15px}}.sheet-wrap{{max-width:100%;overflow:auto}}
.sheet table{{font:12px/1.45 "Segoe UI",sans-serif}}.sheet td{{min-width:64px;max-width:360px;white-space:pre-wrap;overflow-wrap:anywhere}}
.row-head{{position:sticky;left:0;z-index:2;min-width:42px!important;background:#f5f6f8!important;color:var(--muted);text-align:right}}
.slides{{display:grid;gap:24px;justify-items:center}}.slide-wrap{{width:min(100%,1040px)}}.slide-label{{margin:0 0 7px;color:var(--muted);font-size:12px}}
.slide{{position:relative;width:100%;overflow:hidden;background:#fff;box-shadow:0 5px 24px #1f29372b}}.shape{{position:absolute;overflow:hidden}}
.shape-text{{display:flex;white-space:pre-wrap;overflow-wrap:anywhere;line-height:1.2}}.shape img{{display:block;width:100%;height:100%;object-fit:contain}}
.shape table{{width:100%;height:100%;table-layout:fixed;font-size:10px;background:#fff}}.shape td{{padding:2px 4px;overflow:hidden}}
{extra_style}
</style></head><body><div class="preview-meta">{escaped(kind)} · 隔离预览，不执行文档中的脚本或外部资源</div>{body}</body></html>"""


def run_html(run: Any) -> str:
    styles: list[str] = []
    if run.bold:
        styles.append("font-weight:700")
    if run.italic:
        styles.append("font-style:italic")
    if run.underline:
        styles.append("text-decoration:underline")
    if run.font.size is not None:
        styles.append(f"font-size:{run.font.size.pt:.2f}pt")
    color = color_value(run.font.color)
    if color:
        styles.append(f"color:{color}")
    parts = [escaped(run.text).replace("\n", "<br>")]
    try:
        relationships = run.part.related_parts
        for blip in run._r.xpath(".//a:blip"):
            relation_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = relationships.get(relation_id)
            if part is not None and hasattr(part, "blob"):
                content_type = getattr(part, "content_type", "image/png")
                parts.append(f'<img src="{data_url(content_type, part.blob)}" alt="文档图片">')
    except (AttributeError, KeyError, TypeError, ValueError):
        pass
    return f'<span style="{";".join(styles)}">{"".join(parts)}</span>'


def paragraph_html(paragraph: Any) -> str:
    style_name = (paragraph.style.name if paragraph.style else "").lower()
    tag = "p"
    if "heading 1" in style_name or "标题 1" in style_name:
        tag = "h1"
    elif "heading 2" in style_name or "标题 2" in style_name:
        tag = "h2"
    elif "heading 3" in style_name or "标题 3" in style_name:
        tag = "h3"
    content = "".join(run_html(run) for run in paragraph.runs) or escaped(paragraph.text)
    extra = ' class="page-break"' if "page break" in style_name else ""
    if extra:
        return f"<hr{extra}>"
    return f"<{tag}>{content}</{tag}>"


def word_table_html(table: Any) -> str:
    rows = []
    for row in table.rows:
        cells = "".join(f"<td>{''.join(paragraph_html(p) for p in cell.paragraphs)}</td>" for cell in row.cells)
        rows.append(f"<tr>{cells}</tr>")
    return f"<table><tbody>{''.join(rows)}</tbody></table>"


def render_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            parts.append(paragraph_html(Paragraph(child, document)))
        elif child.tag.endswith("}tbl"):
            parts.append(word_table_html(Table(child, document)))
    return document_shell(path.name, "Word 文档", f'<main class="paper">{"".join(parts)}</main>')


def cell_css(cell: Any) -> str:
    styles: list[str] = []
    fill = color_value(cell.fill.fgColor)
    font = color_value(cell.font.color)
    if fill and fill.lower() not in {"#000000", "#ffffff"}:
        styles.append(f"background:{fill}")
    if font:
        styles.append(f"color:{font}")
    if cell.font.bold:
        styles.append("font-weight:700")
    if cell.font.italic:
        styles.append("font-style:italic")
    if cell.alignment.horizontal in {"center", "right"}:
        styles.append(f"text-align:{cell.alignment.horizontal}")
    if cell.alignment.vertical in {"center", "top", "bottom"}:
        styles.append(f"vertical-align:{'middle' if cell.alignment.vertical == 'center' else cell.alignment.vertical}")
    return ";".join(styles)


def cell_text(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return escaped(value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat())
    return escaped(value).replace("\n", "<br>")


def render_xlsx(path: Path) -> str:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(str(path), read_only=False, data_only=False)
    sections: list[str] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets[:MAX_SHEETS], start=1):
            max_row = min(max(worksheet.max_row, 1), MAX_ROWS)
            max_column = min(max(worksheet.max_column, 1), MAX_COLUMNS)
            merged_starts: dict[tuple[int, int], tuple[int, int]] = {}
            merged_covered: set[tuple[int, int]] = set()
            for area in worksheet.merged_cells.ranges:
                if area.min_row > max_row or area.min_col > max_column:
                    continue
                row_span = min(area.max_row, max_row) - area.min_row + 1
                col_span = min(area.max_col, max_column) - area.min_col + 1
                merged_starts[(area.min_row, area.min_col)] = (row_span, col_span)
                for row in range(area.min_row, min(area.max_row, max_row) + 1):
                    for column in range(area.min_col, min(area.max_col, max_column) + 1):
                        if (row, column) != (area.min_row, area.min_col):
                            merged_covered.add((row, column))
            colgroup = []
            for column in range(1, max_column + 1):
                width = worksheet.column_dimensions[get_column_letter(column)].width or 10
                colgroup.append(f'<col style="width:{min(max(width * 7, 54), 280):.0f}px">')
            rows: list[str] = []
            for row_index in range(1, max_row + 1):
                cells = [f'<th class="row-head">{row_index}</th>']
                for column_index in range(1, max_column + 1):
                    if (row_index, column_index) in merged_covered:
                        continue
                    cell = worksheet.cell(row_index, column_index)
                    if isinstance(cell, MergedCell):
                        continue
                    span = merged_starts.get((row_index, column_index), (1, 1))
                    attributes = []
                    if span[0] > 1:
                        attributes.append(f'rowspan="{span[0]}"')
                    if span[1] > 1:
                        attributes.append(f'colspan="{span[1]}"')
                    css = cell_css(cell)
                    if css:
                        attributes.append(f'style="{css}"')
                    cells.append(f'<td {" ".join(attributes)}>{cell_text(cell.value)}</td>')
                rows.append(f"<tr>{''.join(cells)}</tr>")
            truncated = worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS
            note = '<p class="note">工作表较大，预览仅显示前 240 行 × 60 列；完整内容请使用“打开”。</p>' if truncated else ""
            state = "" if worksheet.sheet_state == "visible" else f" · {escaped(worksheet.sheet_state)}"
            sections.append(f'<section class="sheet"><h2>{sheet_index}. {escaped(worksheet.title)}{state}</h2>{note}<div class="sheet-wrap"><table><colgroup><col style="width:42px">{"".join(colgroup)}</colgroup><tbody>{"".join(rows)}</tbody></table></div></section>')
    finally:
        workbook.close()
    if len(workbook.sheetnames) > MAX_SHEETS:
        sections.append('<p class="note">工作簿包含较多工作表，预览仅显示前 12 个。</p>')
    return document_shell(path.name, "Excel 工作簿", f'<main>{"".join(sections)}</main>')


def ppt_fill(shape: Any) -> str | None:
    try:
        if shape.fill.type is None:
            return None
        return color_value(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def ppt_text(shape: Any) -> tuple[str, str]:
    paragraphs: list[str] = []
    styles: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            run_styles = []
            if run.font.bold:
                run_styles.append("font-weight:700")
            if run.font.italic:
                run_styles.append("font-style:italic")
            color = color_value(run.font.color)
            if color:
                run_styles.append(f"color:{color}")
            if run.font.size is not None:
                run_styles.append(f"font-size:{run.font.size.pt:.2f}pt")
            runs.append(f'<span style="{";".join(run_styles)}">{escaped(run.text)}</span>')
        paragraphs.append("".join(runs) or escaped(paragraph.text))
    first_run = next((run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text), None)
    if first_run is not None and first_run.font.size is not None:
        styles.append(f"font-size:{first_run.font.size.pt:.2f}pt")
    anchor = str(getattr(shape.text_frame, "vertical_anchor", "")).lower()
    styles.append("align-items:center" if "middle" in anchor else "align-items:flex-end" if "bottom" in anchor else "align-items:flex-start")
    return "<br>".join(paragraphs), ";".join(styles)


def ppt_table(shape: Any) -> str:
    rows = []
    for row in shape.table.rows:
        rows.append("<tr>" + "".join(f"<td>{escaped(cell.text).replace(chr(10), '<br>')}</td>" for cell in row.cells) + "</tr>")
    return f"<table><tbody>{''.join(rows)}</tbody></table>"


def shape_html(shape: Any, slide_width: int, slide_height: int, z_index: int, image_styles: dict[str, str]) -> str:
    left = 100 * shape.left / slide_width
    top = 100 * shape.top / slide_height
    width = 100 * shape.width / slide_width
    height = 100 * shape.height / slide_height
    rotation = float(getattr(shape, "rotation", 0) or 0)
    base = [
        f"left:{left:.4f}%", f"top:{top:.4f}%", f"width:{width:.4f}%", f"height:{height:.4f}%",
        f"z-index:{z_index}", f"transform:rotate({rotation:.2f}deg)",
    ]
    fill = ppt_fill(shape)
    if fill:
        base.append(f"background:{fill}")
    content = ""
    extra_class = ""
    try:
        if hasattr(shape, "image"):
            blob = shape.image.blob
            class_name = f"ppt-image-{hashlib.sha256(blob).hexdigest()[:16]}"
            image_styles.setdefault(class_name, data_url(shape.image.content_type, blob))
            extra_class = f" {class_name}"
        elif getattr(shape, "has_table", False):
            content = ppt_table(shape)
        elif getattr(shape, "has_text_frame", False) and shape.text_frame.text:
            content, text_style = ppt_text(shape)
            base.append(text_style)
            extra_class = " shape-text"
    except (AttributeError, TypeError, ValueError):
        content = ""
    return f'<div class="shape{extra_class}" style="{";".join(filter(None, base))}">{content}</div>'


def render_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    ratio = presentation.slide_width / presentation.slide_height
    slides: list[str] = []
    image_styles: dict[str, str] = {}
    for index, slide in enumerate(presentation.slides, start=1):
        if index > MAX_SLIDES:
            break
        shapes = "".join(shape_html(shape, presentation.slide_width, presentation.slide_height, z, image_styles) for z, shape in enumerate(slide.shapes, start=1))
        slides.append(f'<section class="slide-wrap"><p class="slide-label">第 {index} 页</p><div class="slide" style="aspect-ratio:{ratio:.8f}">{shapes}</div></section>')
    if len(presentation.slides) > MAX_SLIDES:
        slides.append('<p class="note">演示文稿较长，预览仅显示前 80 页。</p>')
    picture_css = "".join(f'.{name}{{background-image:url("{source}");background-size:contain;background-repeat:no-repeat;background-position:center}}' for name, source in image_styles.items())
    return document_shell(path.name, "PowerPoint 演示文稿", f'<main class="slides">{"".join(slides)}</main>', picture_css)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render a bounded, static Office preview as HTML.")
    parser.add_argument("--input", required=True)
    return parser.parse_args()


def main() -> int:
    path = Path(parse_args().input).resolve(strict=True)
    suffix = path.suffix.lower()
    renderers = {".docx": render_docx, ".xlsx": render_xlsx, ".pptx": render_pptx}
    renderer = renderers.get(suffix)
    if renderer is None:
        raise ValueError(f"unsupported Office preview type: {suffix}")
    sys.stdout.write(renderer(path))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        raise SystemExit(2)
